"""
Conferência dos dados carregados pelo ETL contra as planilhas de origem.

Lê as fontes (extratos TabNet em xlsx e CSVs de empenho da LAI) com um leitor
próprio, que não reaproveita nenhuma função de `etl/`, e compara com o que os
loaders gravam. A comparação cobre todos os valores, não uma amostra.

Duas comparações possíveis:

  1. Fonte × saída do ETL (padrão, não precisa do Neo4j). Os loaders rodam com
     uma sessão falsa que só registra as linhas que seriam gravadas, e a
     agregação DespesaAnual (feita em Cypher no loader) é refeita em Python.
  2. Fonte × grafo (`--neo4j`). Lê IndicadorSaude, as quebras dimensionais e
     DespesaAnual do banco e compara com as mesmas fontes.

Saída em `--out`: resumo.json, divergencias.csv e amostra.csv (valores
sorteados com a célula de origem, para conferência manual no TabNet).

Uso (na raiz do backend):

    python -m scripts.conferencia_etl --out ../conferencia_etl
    python -m scripts.conferencia_etl --neo4j --out ../conferencia_etl
"""

from __future__ import annotations

import argparse
import csv
import json
import math
import os
import random
import re
import sys
import unicodedata
from collections import defaultdict
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

DADOS = Path(os.environ.get("DADOS_DIR", Path(__file__).resolve().parents[2] / "Dados"))
DATASUS = DADOS / "Sorocaba_DATASUS_2015-2025"
ORCAMENTO = DADOS / "orcamento"
ANO_MIN, ANO_MAX = 2016, 2025
TOL = 1e-6

# --- Especificação do mapeamento fonte → grafo -----------------------------
# É a regra de negócio (qual aba vira qual indicador). A leitura das células
# é feita aqui, de forma independente do loader.

SIM, SIH, SINASC = (
    "Sorocaba_Mortalidade_SIM_2015-2025.xlsx",
    "Sorocaba_Internacoes_SIH_2015-2025.xlsx",
    "Sorocaba_Nascidos_Vivos_SINASC_2015-2025.xlsx",
)
SINAN = "Sorocaba_SINAN_Doencas_Notificacao_2015-2025.xlsx"
CNES_AD = "Sorocaba_CNES_Dados_Adicionais_2015-2025.xlsx"

# (arquivo, aba, sistema, subtipo, relação, chave da dimensão, só dezembro)
# chave da dimensão: "escopo" = "sistema:nome", "cid" = código 01-21, "nome" = nome
TABELAS_LARGAS = [
    (SIM, "Óbitos por Faixa Etária", "sim", "mortalidade", "POR_FAIXA_ETARIA", "escopo", False),
    (SIM, "Óbitos por Capítulo CID-10", "sim", "mortalidade", "POR_CAPITULO_CID", "cid", False),
    (SIH, "Internações por Capítulo CID-10", "sih", "internacoes", "POR_CAPITULO_CID", "cid", False),
    (SINASC, "Nascimentos por Idade da Mãe", "sinasc", "nascidos_vivos", "POR_FAIXA_ETARIA", "escopo", False),
    (SINASC, "Nascimentos por Peso ao Nascer", "sinasc", "nascidos_vivos", "POR_FAIXA_PESO", "nome", False),
    (SINAN, "Dengue por Faixa Etária", "sinan", "dengue", "POR_FAIXA_ETARIA", "escopo", False),
    (SINAN, "Chikungunya por Faixa Etária", "sinan", "chikungunya", "POR_FAIXA_ETARIA", "escopo", False),
    (SINAN, "Sífilis Adquirida", "sinan", "sifilis_adquirida", "POR_FAIXA_ETARIA", "escopo", False),
    (SINAN, "Sífilis em Gestante", "sinan", "sifilis_gestante", "POR_FAIXA_ETARIA", "escopo", False),
    (SINAN, "Sífilis Congênita", "sinan", "sifilis_congenita", "POR_FAIXA_ETARIA", "escopo", False),
    (SINAN, "Coqueluche por Faixa Etária", "sinan", "coqueluche", "POR_FAIXA_ETARIA", "escopo", False),
    (SINAN, "Hepatites Virais", "sinan", "hepatites_virais", "POR_FAIXA_ETARIA", "escopo", False),
    (SINAN, "Tuberculose por Sexo", "sinan", "tuberculose", "POR_SEXO", "nome", False),
    (SINAN, "Hanseníase por Sexo", "sinan", "hanseniase", "POR_SEXO", "nome", False),
    ("Sorocaba_Rede_Assistencial_CNES_2015-2025.xlsx", "Estabelecimentos por Tipo", "cnes",
     "estabelecimentos_por_tipo", "POR_TIPO_ESTABELECIMENTO", "nome", False),
    (CNES_AD, "Rec.Humanos-Ocupacoes(Dez)", "cnes", "ocupacoes", "POR_OCUPACAO", "nome", False),
    (CNES_AD, "Equipes de Saude", "cnes", "equipes_saude", "POR_TIPO_EQUIPE", "nome", True),
    (CNES_AD, "Tipo Atendimento (SUS)", "cnes", "tipo_atendimento", "POR_TIPO_ATENDIMENTO", "nome", True),
    (CNES_AD, "Rec.Fisicos-Leitos e Consult", "cnes", "leitos_consultorios", "POR_TIPO_LEITO_CONSULTORIO", "nome", True),
    (CNES_AD, "Rec.Fisicos-Equipamentos", "cnes", "equipamentos", "POR_TIPO_EQUIPAMENTO", "nome", True),
]
VALOR_UNICO = [
    ("Sorocaba_CNES_Recursos_2015-2025.xlsx", "Leitos de Internacao (Dez)", "cnes", "leitos"),
    ("Sorocaba_CNES_Recursos_2015-2025.xlsx", "Profissionais de Saude (Dez)", "cnes", "profissionais"),
    ("Sorocaba_SIA_Producao_Ambulatorial_2015-2025.xlsx", "Producao Ambulatorial por Ano", "sia", "producao_ambulatorial"),
]
METRICAS_ESTABELECIMENTOS = {
    "Nivel_Atencao_Qtd_Geral": "estabelecimentos_nivel_atencao",
    "Servico_Classificacao_a_partir_Mar2008": "estabelecimentos_servico_classificacao",
    "Habilitacao_Qtd_Estab_Habilitados": "estabelecimentos_habilitacao",
    "Vigilancia_Epidem_Sanitaria_Qtd": "estabelecimentos_vigilancia_epidemiologica",
}
SIPNI = "Sorocaba_SIPNI_Cobertura_Vacinal_2015-2026.xlsx"
VACINAS = [  # (abas, subtipo, relação, agregação do total anual)
    (("Cobertura 2015-2022", "Cobertura 2023-2025"), "cobertura_vacinal", "COBERTURA_VACINAL", "media"),
    (("Doses 2015-2022", "Doses 2023-2026"), "doses_aplicadas", "DOSES_APLICADAS", "soma"),
]
COVID = "Sorocaba_COVID19_Casos_Obitos_2020-2023.xlsx"

SUBFUNCOES = {
    "ATENCAO BASICA": 301, "ASSISTENCIA HOSPITALAR E AMBULATORIAL": 302,
    "SUPORTE PROFILATICO E TERAPEUTICO": 303, "VIGILANCIA SANITARIA": 304,
    "VIGILANCIA EPIDEMIOLOGICA": 305, "ALIMENTACAO E NUTRICAO": 306, "ADMINISTRACAO GERAL": 122,
}
ROMANOS = {"I": 1, "V": 5, "X": 10, "L": 50}
MESES = {"Jan": 1, "Fev": 2, "Mar": 3, "Abr": 4, "Mai": 5, "Jun": 6,
         "Jul": 7, "Ago": 8, "Set": 9, "Out": 10, "Nov": 11, "Dez": 12}


# --- Leitura independente das fontes ----------------------------------------

def _romano(s: str) -> int:
    total, anterior = 0, 0
    for c in reversed(s):
        v = ROMANOS[c]
        total = total - v if v < anterior else total + v
        anterior = max(anterior, v)
    return total


def _codigo_cid(cabecalho: str) -> str | None:
    m = re.match(r"\s*Cap\s*([IVXL]+|\d+)\s*\(", cabecalho)
    if not m:
        return None
    g = m.group(1)
    return f"{int(g) if g.isdigit() else _romano(g):02d}"


def _numero(v):
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    return "invalido"


def _ano_da_linha(chave, mensal: bool) -> int | None:
    texto = str(chave).strip()
    if mensal:
        m = re.fullmatch(r"(\d{4})/(\w{3})", texto)
        return int(m.group(1)) if m and MESES.get(m.group(2)) == 12 else None
    return int(texto) if re.fullmatch(r"\d{4}", texto) else None


def _linhas(arquivo: str, aba: str) -> list[tuple]:
    wb = openpyxl.load_workbook(DATASUS / arquivo, read_only=True, data_only=True)
    try:
        return [tuple(r) for r in wb[aba].iter_rows(values_only=True)]
    finally:
        wb.close()


class Esperado:
    """Valores que o grafo deveria conter, com a célula de origem de cada um."""

    def __init__(self) -> None:
        self.indicadores: dict[str, list[dict]] = defaultdict(list)   # chave -> ocorrências na fonte
        self.quebras: dict[tuple, dict] = {}                           # (chave, relação, dimensão) -> valor
        self.celulas_invalidas: list[dict] = []

    def indicador(self, sistema, subtipo, ano, mes, valor, origem, acumulado=None):
        chave = f"{sistema}:{subtipo}:{ano}:{mes or 0}"
        self.indicadores[chave].append({"valor": valor, "acumulado": acumulado, "origem": origem})
        return chave

    def quebra(self, chave, relacao, dimensao, valor, origem):
        self.quebras[(chave, relacao, dimensao)] = {"valor": valor, "origem": origem}


def ler_fontes_saude() -> Esperado:
    esp = Esperado()

    for arquivo, aba, sistema, subtipo, relacao, modo, mensal in TABELAS_LARGAS:
        linhas = _linhas(arquivo, aba)
        cab = [str(h).strip() if h is not None else "" for h in linhas[2]]
        col_total = next((i for i, h in enumerate(cab) if i > 0 and h.lower() == "total"), None)
        for r, linha in enumerate(linhas[3:], start=4):
            ano = _ano_da_linha(linha[0], mensal) if linha and linha[0] is not None else None
            if ano is None or not (ANO_MIN <= ano <= ANO_MAX):
                continue
            total = _numero(linha[col_total]) if col_total is not None else None
            origem = f"{arquivo} | {aba} | {get_column_letter(col_total + 1)}{r}" if col_total is not None else f"{arquivo} | {aba} | (sem coluna Total)"
            if total == "invalido":
                esp.celulas_invalidas.append({"origem": origem, "valor": linha[col_total]})
                total = None
            chave = esp.indicador(sistema, subtipo, ano, None, total, origem)
            for c in range(1, len(cab)):
                if c == col_total or not cab[c]:
                    continue
                v = _numero(linha[c]) if c < len(linha) else None
                cel = f"{arquivo} | {aba} | {get_column_letter(c + 1)}{r}"
                if v == "invalido":
                    esp.celulas_invalidas.append({"origem": cel, "valor": linha[c]})
                    continue
                if v is None:
                    continue
                if modo == "cid":
                    dim = _codigo_cid(cab[c])
                    if dim is None:
                        continue
                elif modo == "escopo":
                    dim = f"{sistema}:{cab[c]}"
                else:
                    dim = cab[c]
                esp.quebra(chave, relacao, dim, v, cel)

    for arquivo, aba, sistema, subtipo in VALOR_UNICO:
        for r, linha in enumerate(_linhas(arquivo, aba)[3:], start=4):
            ano = _ano_da_linha(linha[0], False) if linha and linha[0] is not None else None
            v = _numero(linha[1]) if ano and len(linha) > 1 else None
            if ano is None or not (ANO_MIN <= ano <= ANO_MAX) or v is None:
                continue
            esp.indicador(sistema, subtipo, ano, None, v, f"{arquivo} | {aba} | B{r}")

    linhas = _linhas(CNES_AD, "Estabelecimentos")
    cab = [str(h).strip() if h is not None else "" for h in linhas[2]]
    for r, linha in enumerate(linhas[3:], start=4):
        ano = _ano_da_linha(linha[0], True) if linha and linha[0] is not None else None
        if ano is None or not (ANO_MIN <= ano <= ANO_MAX):
            continue
        for c, nome in enumerate(cab):
            subtipo = METRICAS_ESTABELECIMENTOS.get(nome)
            v = _numero(linha[c]) if subtipo and c < len(linha) else None
            if subtipo and isinstance(v, float):
                esp.indicador("cnes", subtipo, ano, None, v, f"{CNES_AD} | Estabelecimentos | {get_column_letter(c + 1)}{r}")

    for abas, subtipo, relacao, agregacao in VACINAS:
        por_ano: dict[int, list[tuple[str, float, str]]] = defaultdict(list)
        for aba in abas:
            linhas = _linhas(SIPNI, aba)
            for c, h in enumerate(linhas[2]):
                ano = int(h) if re.fullmatch(r"\d{4}", str(h).strip()) else None
                if ano is None or not (ANO_MIN <= ano <= ANO_MAX):
                    continue
                for r, linha in enumerate(linhas[3:], start=4):
                    vacina = linha[0]
                    v = _numero(linha[c]) if c < len(linha) else None
                    if vacina is None or not isinstance(v, float):
                        continue
                    por_ano[ano].append((str(vacina).strip(), v, f"{SIPNI} | {aba} | {get_column_letter(c + 1)}{r}"))
        for ano, itens in por_ano.items():
            valores = [v for _, v, _ in itens]
            total = sum(valores) if agregacao == "soma" else sum(valores) / len(valores)
            chave = esp.indicador("sipni", subtipo, ano, None, total,
                                  f"{SIPNI} | {agregacao} de {len(itens)} vacinas em {ano}")
            for vacina, v, cel in itens:
                esp.quebra(chave, relacao, vacina, v, cel)

    for aba, mensal in (("COVID Casos-Obitos Anual", False), ("COVID Casos-Obitos Mensal", True)):
        for r, linha in enumerate(_linhas(COVID, aba)[3:], start=4):
            texto = str(linha[0]).strip() if linha and linha[0] is not None else ""
            m = re.fullmatch(r"(\d{4})-(\d{2})", texto) if mensal else re.fullmatch(r"(\d{4})", texto)
            if not m:
                continue
            ano, mes = int(m.group(1)), (int(m.group(2)) if mensal else None)
            if not (ANO_MIN <= ano <= ANO_MAX):
                continue
            for subtipo, c_novo, c_acum in (("casos", 1, 3), ("obitos", 2, 4)):
                novo, acum = _numero(linha[c_novo]), _numero(linha[c_acum])
                esp.indicador("covid", subtipo, ano, mes, novo if isinstance(novo, float) else None,
                              f"{COVID} | {aba} | {get_column_letter(c_novo + 1)}{r}",
                              acum if isinstance(acum, float) else None)
    return esp


def _normaliza(texto: str) -> str:
    sem_acento = unicodedata.normalize("NFKD", str(texto)).encode("ascii", "ignore").decode()
    return sem_acento.strip().upper()


def ler_fontes_orcamento() -> tuple[dict, dict]:
    """Soma os empenhos por (subfunção, ano) a partir dos CSVs brutos."""
    somas = defaultdict(lambda: {"valorProcessado": 0.0, "valorPago": 0.0, "valorEmpenhado": 0.0, "linhas": 0})
    descartes = defaultdict(int)
    for csv_path in sorted(ORCAMENTO.glob("*/*.csv")):
        nome = csv_path.name.lower()
        if "previsto_e_empenhado" in nome or "empenhado" not in nome:
            continue
        with open(csv_path, encoding="latin-1", newline="") as f:
            for row in csv.DictReader(f, delimiter="|"):
                try:
                    data = datetime.strptime(row["Data Empenho"].strip(), "%d/%m/%Y")
                except (ValueError, AttributeError):
                    descartes["data_invalida"] += 1
                    continue
                if not (ANO_MIN <= data.year <= ANO_MAX):
                    descartes[f"fora_do_periodo_{data.year}"] += 1
                    continue
                codigo = SUBFUNCOES.get(_normaliza(row["Sub Função"]))
                if codigo is None:
                    descartes["subfuncao_nao_mapeada"] += 1
                    continue
                s = somas[(codigo, data.year)]
                try:
                    s["valorProcessado"] += float(row["Valor Processado"] or 0)
                    s["valorPago"] += float(row["Valor Pago"] or 0)
                    s["valorEmpenhado"] += float(row["Saldo Empenho"] or 0)
                except ValueError:
                    descartes["valor_nao_numerico"] += 1
                s["linhas"] += 1
    return dict(somas), dict(descartes)


# --- O que o ETL grava -------------------------------------------------------

class _SessaoQueRegistra:
    """Sessão Neo4j falsa: guarda as linhas de cada `session.run`, na ordem."""

    def __init__(self) -> None:
        self.indicadores: list[dict] = []
        self.quebras: list[tuple[str, dict]] = []

    def run(self, query: str, **params):
        rows = params.get("rows")
        if rows is None:
            return None
        if "MERGE (i:IndicadorSaude" in query:
            self.indicadores.extend(rows)
        else:
            m = re.search(r"MERGE \(i\)-\[r:(\w+)\]->\(d\)", query)
            if m:
                self.quebras.extend((m.group(1), row) for row in rows)
        return None


def saida_do_etl() -> tuple[dict, dict, dict, dict]:
    from etl import orcamento_loader, saude_indicadores_loader as sl

    sessao = _SessaoQueRegistra()
    for _, fn in sl._LOADERS:
        fn(sessao, "conferencia")

    gravacoes = defaultdict(list)  # chave -> valores na ordem de gravação (MERGE + SET: vale o último)
    for row in sessao.indicadores:
        gravacoes[row["chave"]].append((row["valorTotal"], row.get("valorAcumulado")))
    quebras = {(row["indicadorChave"], rel, row["chave"]): row["valor"] for rel, row in sessao.quebras}

    despesas = defaultdict(lambda: {"valorProcessado": 0.0, "valorPago": 0.0, "valorEmpenhado": 0.0, "linhas": 0})
    for periodo in sorted(p for p in orcamento_loader.ORCAMENTO_DIR.iterdir() if p.is_dir()):
        for csv_path in sorted(periodo.glob("*.csv")):
            nome = csv_path.name.lower()
            if "previsto_e_empenhado" in nome or "empenhado" not in nome:
                continue
            for e in orcamento_loader._read_empenho_csv(csv_path, periodo.name):
                d = despesas[(e["subfuncaoCodigo"], e["ano"])]
                d["valorProcessado"] += e["valorProcessado"]
                d["valorPago"] += e["valorPago"]
                d["valorEmpenhado"] += e["saldoEmpenho"]
                d["linhas"] += 1
    return dict(gravacoes), quebras, dict(despesas), {}


def saida_do_grafo() -> tuple[dict, dict, dict, dict]:
    from dotenv import load_dotenv
    from neo4j import GraphDatabase

    load_dotenv()
    uri = os.environ.get("NEO4J_URI", "bolt://localhost:7687").replace("bolt://neo4j:", "bolt://localhost:")
    driver = GraphDatabase.driver(uri, auth=(os.environ["NEO4J_USER"], os.environ["NEO4J_PASSWORD"]))
    with driver.session() as s:
        gravacoes = {r["chave"]: [(r["v"], r["a"])] for r in s.run(
            "MATCH (i:IndicadorSaude) RETURN i.chave AS chave, i.valorTotal AS v, i.valorAcumulado AS a")}
        quebras = {(r["chave"], r["rel"], r["dim"]): r["valor"] for r in s.run(
            "MATCH (i:IndicadorSaude)-[r]->(d) WHERE r.valor IS NOT NULL "
            "RETURN i.chave AS chave, type(r) AS rel, coalesce(d.chave, d.codigo) AS dim, r.valor AS valor")}
        despesas = {(r["sf"], r["ano"]): {"valorProcessado": r["vp"], "valorPago": r["pg"], "valorEmpenhado": r["ve"]}
                    for r in s.run("MATCH (d:DespesaAnual) RETURN d.subfuncaoCodigo AS sf, d.ano AS ano, "
                                   "d.valorProcessado AS vp, d.valorPago AS pg, d.valorEmpenhado AS ve")}
        meta = {"uri": uri, "indicadores": len(gravacoes), "quebras": len(quebras), "despesas_anuais": len(despesas)}
    driver.close()
    return gravacoes, quebras, despesas, meta


# --- Comparação ----------------------------------------------------------------

def _igual(a, b) -> bool:
    if a is None or b is None:
        return a is None and b is None
    return math.isclose(float(a), float(b), rel_tol=1e-9, abs_tol=TOL)


def comparar(esp: Esperado, orc: dict, gravacoes: dict, quebras: dict, despesas: dict) -> dict:
    div: list[dict] = []
    res = {"indicadores": {"conferidos": 0, "iguais": 0}, "quebras": {"conferidos": 0, "iguais": 0},
           "despesas": {"conferidos": 0, "iguais": 0}}

    for chave, ocorrencias in esp.indicadores.items():
        res["indicadores"]["conferidos"] += 1
        gravado = gravacoes.get(chave)
        final = gravado[-1] if gravado else None
        if len({o["valor"] for o in ocorrencias}) > 1:
            div.append({"tipo": "fonte_com_totais_diferentes", "chave": chave,
                        "fonte": "; ".join(f"{o['valor']} ({o['origem']})" for o in ocorrencias),
                        "gravado": final[0] if final else None})
        alvo = ocorrencias[-1]  # o loader grava as abas na ordem da especificação; vale a última
        if final is None:
            div.append({"tipo": "indicador_ausente", "chave": chave, "fonte": alvo["valor"], "origem": alvo["origem"]})
        elif _igual(final[0], alvo["valor"]) and _igual(final[1], alvo["acumulado"]):
            res["indicadores"]["iguais"] += 1
        else:
            div.append({"tipo": "indicador_diferente", "chave": chave, "fonte": alvo["valor"],
                        "gravado": final[0], "origem": alvo["origem"]})
    extras = sorted(set(gravacoes) - set(esp.indicadores))
    div += [{"tipo": "indicador_sem_origem", "chave": c, "gravado": gravacoes[c][-1][0]} for c in extras]

    for chave, item in esp.quebras.items():
        res["quebras"]["conferidos"] += 1
        g = quebras.get(chave)
        if _igual(g, item["valor"]):
            res["quebras"]["iguais"] += 1
        else:
            div.append({"tipo": "quebra_ausente" if g is None else "quebra_diferente", "chave": " | ".join(chave),
                        "fonte": item["valor"], "gravado": g, "origem": item["origem"]})
    div += [{"tipo": "quebra_sem_origem", "chave": " | ".join(k), "gravado": quebras[k]}
            for k in sorted(set(quebras) - set(esp.quebras))]

    for chave, soma in orc.items():
        res["despesas"]["conferidos"] += 1
        g = despesas.get(chave)
        campos = ("valorProcessado", "valorPago", "valorEmpenhado")
        if g and all(abs(g[c] - soma[c]) < 0.005 for c in campos):
            res["despesas"]["iguais"] += 1
        else:
            div.append({"tipo": "despesa_ausente" if g is None else "despesa_diferente",
                        "chave": f"subfuncao {chave[0]} | {chave[1]}",
                        "fonte": round(soma["valorProcessado"], 2),
                        "gravado": round(g["valorProcessado"], 2) if g else None})
    div += [{"tipo": "despesa_sem_origem", "chave": f"subfuncao {k[0]} | {k[1]}"} for k in sorted(set(despesas) - set(orc))]

    sobrescritas = [{"chave": c, "gravacoes": [v for v, _ in vs]} for c, vs in gravacoes.items()
                    if len(vs) > 1 and len({v for v, _ in vs}) > 1]
    res["divergencias"] = len(div)
    res["divergencias_por_tipo"] = dict(sorted(((t, sum(1 for d in div if d["tipo"] == t)) for t in {d["tipo"] for d in div})))
    res["sobrescritas_com_valor_diferente"] = sobrescritas
    return res, div


def consistencia_interna(esp: Esperado) -> list[dict]:
    """Total da planilha × soma das categorias (checagem da própria fonte)."""
    somas = defaultdict(float)
    tem = defaultdict(bool)
    for (chave, rel, _), item in esp.quebras.items():
        somas[(chave, rel)] += item["valor"]
        tem[(chave, rel)] = True
    saida = []
    for (chave, rel), soma in somas.items():
        if rel in ("COBERTURA_VACINAL", "POR_TIPO_ATENDIMENTO", "POR_TIPO_LEITO_CONSULTORIO"):
            continue  # não são partes de um todo: o total não é a soma das categorias
        for o in esp.indicadores.get(chave, []):
            if o["valor"] is not None and not o["origem"].startswith(SIPNI) and rel_da_origem(o["origem"], rel):
                if abs(o["valor"] - soma) > TOL:
                    saida.append({"chave": chave, "relacao": rel, "total_da_planilha": o["valor"],
                                  "soma_das_categorias": soma, "diferenca": soma - o["valor"], "origem": o["origem"]})
    return saida


def rel_da_origem(origem: str, rel: str) -> bool:
    aba = origem.split(" | ")[1] if " | " in origem else ""
    return any(a == aba and r == rel for _, a, _, _, r, _, _ in TABELAS_LARGAS)


def amostra(esp: Esperado, gravacoes: dict, orc: dict, despesas: dict, n: int, seed: int) -> list[dict]:
    """Amostra estratificada por sistema, para que toda fonte apareça na conferência manual."""
    rnd = random.Random(seed)
    por_sistema = defaultdict(list)
    for c, o in sorted(esp.indicadores.items()):
        if o[-1]["valor"] is not None:
            por_sistema[c.split(":")[0]].append(c)
    for chaves in por_sistema.values():
        rnd.shuffle(chaves)
    sorteadas = []
    while len(sorteadas) < n and any(por_sistema.values()):
        for sistema in sorted(por_sistema):
            if por_sistema[sistema] and len(sorteadas) < n:
                sorteadas.append(por_sistema[sistema].pop())
    linhas = []
    for c in sorteadas:
        o = esp.indicadores[c][-1]
        g = gravacoes.get(c)
        linhas.append({"tipo": "indicador", "chave": c, "origem": o["origem"], "valor_na_fonte": o["valor"],
                       "valor_gravado": g[-1][0] if g else None})
    for k in rnd.sample(sorted(orc), min(max(n // 3, 1), len(orc))):
        linhas.append({"tipo": "despesa_anual", "chave": f"subfuncao {k[0]} | {k[1]}",
                       "origem": f"Dados/orcamento (soma de {orc[k]['linhas']} empenhos)",
                       "valor_na_fonte": round(orc[k]["valorProcessado"], 2),
                       "valor_gravado": round(despesas[k]["valorProcessado"], 2) if k in despesas else None})
    return linhas


def _grava_csv(path: Path, linhas: list[dict]) -> None:
    campos = sorted({k for l in linhas for k in l}) or ["vazio"]
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=campos, delimiter=";")
        w.writeheader()
        w.writerows(linhas)


def main() -> int:
    sys.stdout.reconfigure(encoding="utf-8")
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--neo4j", action="store_true", help="compara a fonte com o grafo, não com a saída do ETL")
    ap.add_argument("--out", default="conferencia_etl", help="pasta de saída")
    ap.add_argument("--amostra", type=int, default=30, help="tamanho da amostra para conferência manual")
    ap.add_argument("--seed", type=int, default=42)
    args = ap.parse_args()

    esp = ler_fontes_saude()
    orc, descartes = ler_fontes_orcamento()
    alvo = "grafo Neo4j" if args.neo4j else "saída do ETL (sem banco)"
    gravacoes, quebras, despesas, meta = saida_do_grafo() if args.neo4j else saida_do_etl()

    res, div = comparar(esp, orc, gravacoes, quebras, despesas)
    res["comparado_com"] = alvo
    res["meta"] = meta
    res["orcamento_descartes_na_fonte"] = descartes
    res["orcamento_empenhos_considerados"] = sum(s["linhas"] for s in orc.values())
    res["celulas_nao_numericas_na_fonte"] = esp.celulas_invalidas
    res["fonte_total_diferente_da_soma"] = consistencia_interna(esp)

    out = Path(args.out)
    out.mkdir(parents=True, exist_ok=True)
    (out / "resumo.json").write_text(json.dumps(res, ensure_ascii=False, indent=2), encoding="utf-8")
    _grava_csv(out / "divergencias.csv", div)
    _grava_csv(out / "amostra.csv", amostra(esp, gravacoes, orc, despesas, args.amostra, args.seed))

    print(f"Comparado com: {alvo}")
    for k in ("indicadores", "quebras", "despesas"):
        print(f"  {k}: {res[k]['iguais']}/{res[k]['conferidos']} iguais")
    print(f"  divergências: {res['divergencias']} {res['divergencias_por_tipo']}")
    print(f"  sobrescritas com valor diferente: {len(res['sobrescritas_com_valor_diferente'])}")
    print(f"  total da planilha diferente da soma das categorias: {len(res['fonte_total_diferente_da_soma'])}")
    print(f"  células não numéricas na fonte: {len(esp.celulas_invalidas)}")
    print(f"  empenhos considerados: {res['orcamento_empenhos_considerados']} | descartados: {descartes}")
    print(f"Arquivos em {out.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
