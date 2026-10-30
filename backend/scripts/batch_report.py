"""
Gera o relatório XLSX (tabelas, gráficos e comparação estrela vs.
hierárquica) a partir das linhas produzidas por `scripts.batch_eval`.

As abas de resumo usam fórmulas (AVERAGEIFS/COUNTIFS/SUMIFS) sobre a aba
`Resultados`, então editar/filtrar os dados brutos recalcula tudo no
Excel. O arquivo é marcado para recalcular ao abrir (visualizadores que
não calculam fórmulas, como o preview do VS Code, mostram células vazias).

Uso para regenerar o XLSX de uma execução anterior, sem rodar nada de
novo (o `batch_eval` grava um `.jsonl` ao lado do `.xlsx`):

    python -m scripts.batch_report resultados.jsonl resultados.xlsx
"""

from __future__ import annotations

import json
import sys
from datetime import datetime
from typing import Any

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.chart import BarChart, Reference, Series
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

FONT = "Arial"
COR_ESTRELA = "2F6EBA"
COR_HIER = "E07B39"
COR_NEUTRA = "5B9A6E"

_F_BASE = Font(name=FONT, size=10)
_F_BOLD = Font(name=FONT, size=10, bold=True)
_F_HEAD = Font(name=FONT, size=10, bold=True, color="FFFFFF")
_F_TITLE = Font(name=FONT, size=16, bold=True, color="1F3864")
_F_SUB = Font(name=FONT, size=11, bold=True, color="1F3864")
_F_NOTE = Font(name=FONT, size=9, italic=True, color="595959")
_FILL_HEAD = PatternFill("solid", fgColor="1F3864")
_FILL_ZEBRA = PatternFill("solid", fgColor="F2F5FA")
_BORDER = Border(bottom=Side(style="thin", color="D0D7E2"))

# (chave em `rows`, cabeçalho, largura, number_format)
COLUNAS: list[tuple[str, str, int, str | None]] = [
    ("model", "Modelo", 18, None),
    ("category", "Categoria", 24, None),
    ("scenario", "Cenário", 38, None),
    ("topology", "Topologia", 13, None),
    ("question_len", "Tam. pergunta (chars)", 14, "#,##0"),
    ("reference_year", "Ano de referência", 10, "0"),
    ("intent_success", "Interpretação OK", 13, None),
    ("intent_missing", "Motivo da recusa", 18, None),
    ("intent_elapsed_ms", "Interpretação (ms)", 14, "#,##0"),
    ("intent_tokens", "Interpretação (tokens)", 14, "#,##0"),
    ("health_params_inferred", "Tema inferido", 11, None),
    ("date_range_inferred", "Período inferido", 11, None),
    ("resolved_date_from", "Ano inicial", 10, "0"),
    ("resolved_date_to", "Ano final", 10, "0"),
    ("resolved_health_params_count", "Nº indicadores", 11, "0"),
    ("resolved_health_params", "Indicadores", 36, None),
    ("wall_clock_ms", "Tempo de parede (ms)", 14, "#,##0"),
    ("tempo_total_ms", "Tempo até o texto (ms)", 14, "#,##0"),
    ("total_tokens", "Tokens totais", 12, "#,##0"),
    ("prompt_tokens", "Tokens prompt", 12, "#,##0"),
    ("completion_tokens", "Tokens resposta", 12, "#,##0"),
    ("llm_call_count", "Chamadas LLM", 11, "0"),
    ("overhead_percent", "Overhead coord. (%)", 13, "0.00"),
    ("agent_invocations", "Invocações de agente", 12, "0"),
    ("message_count", "Mensagens", 11, "0"),
    ("mensagens_dados", "Mensagens de dados", 11, "0"),
    ("mensagens_texto", "Mensagens de texto", 11, "0"),
    ("caracteres_texto", "Caracteres em mensagens de texto", 13, "#,##0"),
    ("ferramentas_acionadas", "Ferramentas acionadas", 11, "0"),
    ("texto_analise_len", "Tam. texto (chars)", 12, "#,##0"),
    ("ragas_faithfulness", "RAGAS faithfulness", 12, "0.000"),
    ("ragas_answer_relevancy", "RAGAS answer relevancy", 12, "0.000"),
    ("ragas_context_relevance", "RAGAS context relevance", 12, "0.000"),
    ("ragas_juiz", "Juiz RAGAS", 12, None),
    ("ragas_prompt_tokens", "Juiz: tokens de entrada", 12, "#,##0"),
    ("ragas_completion_tokens", "Juiz: tokens de saída", 12, "#,##0"),
    ("ragas_total_tokens", "Juiz: tokens totais", 12, "#,##0"),
    ("fidelidade_numerica", "Fidelidade numérica", 12, "0.000"),
    ("numeros_texto", "Números no texto", 10, "0"),
    ("numeros_nao_verificados", "Números não verificados", 11, "0"),
    # Planejamento por LLM (agents/planejamento.py) e redação
    ("plano_nos", "Nós que planejaram", 10, "0"),
    ("plano_origem", "Origem do plano", 30, None),
    ("plano_acrescimos", "Plano: acréscimos", 10, "0"),
    ("plano_retiradas", "Plano: retiradas", 10, "0"),
    ("redacao_modo", "Redação", 10, None),
    ("redacao_ramos", "Redação: ramos", 10, "0"),
    ("redacao_falhas", "Redação: ramos que falharam", 11, "0"),
    ("overhead_ms", "Overhead (ms)", 11, "#,##0.0"),
    ("overhead_deliberacao_ms", "Overhead: deliberação LLM (ms)", 13, "#,##0.0"),
    ("overhead_estrutural_ms", "Overhead: estrutura (ms)", 12, "#,##0.0"),
    ("fase_dominio_ms", "Fase domínio (ms)", 12, "#,##0"),
    ("fase_analitico_ms", "Fase analítico (ms)", 12, "#,##0"),
    ("fase_contexto_ms", "Fase contexto (ms)", 12, "#,##0"),
    ("fase_supervisores_ms", "Fase coordenação (ms)", 12, "#,##0.0"),
    ("completude", "Completude (0-1)", 11, "0.000"),
    ("completude_correlacoes", "Completude: correlações", 12, "0.00"),
    ("completude_anomalias", "Completude: anomalias", 12, "0.00"),
    ("completude_contexto", "Completude: contexto", 12, "0.00"),
    ("payload_records", "Registros transportados", 12, "#,##0"),
    ("sucesso", "Análise bem-sucedida (1/0)", 12, "0"),
    ("dentro_do_orcamento_de_tempo", "Dentro do orçamento de tempo (1/0)", 13, "0"),
    ("consistencia_identica", "Correlações/anomalias idênticas (1/0)", 14, "0"),
    ("fault", "Falha injetada", 20, None),
    ("r1_score", "Cobertura R1 (0-1)", 12, "0.00"),
    ("r1_completed", "Componentes completos (de 7)", 13, "0"),
    ("texto_gerado", "Texto gerado (1/0)", 11, "0"),
    ("error_events", "Eventos de erro", 11, "0"),
    # Revisão das entregas (agents/revisao_entregas.py)
    ("revisao_chamadas", "Revisão: chamadas LLM", 11, "0"),
    ("revisao_registros", "Revisão: registros revistos", 12, "#,##0"),
    ("revisao_sinalizados", "Revisão: sinalizados", 11, "0"),
    ("revisao_confirmados", "Revisão: confirmados", 11, "0"),
    ("revisao_corrigidos", "Revisão: corrigidos", 11, "0"),
    ("revisao_descartados", "Revisão: descartados", 11, "0"),
    ("revisao_nao_verificados", "Revisão: não verificados", 11, "0"),
    ("revisao_reenvios", "Revisão: reenvios", 11, "0"),
    ("revisao_falhou", "Revisão falhou (1/0)", 11, "0"),
    # Falha semântica (scripts/falhas_semanticas.py), medida contra o gabarito
    ("alvo_falha", "Agente corrompido", 12, None),
    ("escopo", "Escopo da pergunta", 13, None),
    ("corrupcao_registros", "Registros corrompidos", 12, "0"),
    ("revisao_vp", "Sinais em registro corrompido", 13, "0"),
    ("revisao_fp", "Sinais em registro limpo", 12, "0"),
    ("contaminados_finais", "Corrompidos no resultado final", 13, "0"),
    ("perdidos", "Corrompidos descartados", 12, "0"),
    ("achados_iguais_limpos", "Achados iguais aos dos dados limpos (1/0)", 14, "0"),
    ("intent_reused", "Interpretação reaproveitada", 13, None),
    ("run_error", "Erro na execução", 34, None),
    ("clarification_message", "Mensagem ao usuário", 50, None),
    ("question", "Pergunta", 60, None),
]

# Colunas que só aparecem quando alguma linha as traz (execuções anteriores
# à revisão das entregas e bateria sem falha semântica não as têm).
_COLUNAS_OPCIONAIS = {
    c[0] for c in COLUNAS
    if c[0].startswith("revisao_")
    or c[0] in ("alvo_falha", "escopo", "corrupcao_registros", "contaminados_finais",
                "perdidos", "achados_iguais_limpos")
}

# Comportamento esperado por categoria — premissa do autor do relatório
# (não vem do sistema): usada só para calcular a "aderência" na aba
# Interpretacao. "Analisar" = deve resolver parâmetros e rodar o pipeline;
# "Recusar" = deve recusar; "Conversa" = deve responder sem análise;
# "Misto" = tanto recusar quanto normalizar é aceitável (não pontuado).
ESPERADO = {
    "fora_de_escopo": "Recusar",
    "prompt_injection": "Recusar",
    "chitchat": "Conversa",
    "data_invalida": "Misto",
}
ESPERADO_PADRAO = "Analisar"


def _limpa(v: Any) -> Any:
    if isinstance(v, str):
        return ILLEGAL_CHARACTERS_RE.sub("", v)
    if isinstance(v, (list, tuple, dict)):
        return json.dumps(v, ensure_ascii=False)
    return v


def _cabecalho(ws, row: int, headers: list[str], col0: int = 1) -> None:
    for i, h in enumerate(headers):
        c = ws.cell(row=row, column=col0 + i, value=h)
        c.font = _F_HEAD
        c.fill = _FILL_HEAD
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 32


def _titulo(ws, titulo: str, sub: str | None = None) -> None:
    ws["A1"] = titulo
    ws["A1"].font = _F_TITLE
    if sub:
        ws["A2"] = sub
        ws["A2"].font = _F_NOTE
    ws.sheet_view.showGridLines = False


def _estilo_corpo(ws, r0: int, r1: int, c0: int, c1: int, fmts: dict[int, str] | None = None) -> None:
    for r in range(r0, r1 + 1):
        for c in range(c0, c1 + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = _F_BASE
            cell.border = _BORDER
            if (r - r0) % 2 == 1:
                cell.fill = _FILL_ZEBRA
            if fmts and c in fmts:
                cell.number_format = fmts[c]


def _larguras(ws, larguras: list[int], col0: int = 1) -> None:
    for i, w in enumerate(larguras):
        ws.column_dimensions[get_column_letter(col0 + i)].width = w


def _barras(titulo: str, y_titulo: str, w: float = 18, h: float = 9) -> BarChart:
    ch = BarChart()
    ch.type = "col"
    ch.grouping = "clustered"
    ch.title = titulo
    ch.y_axis.title = y_titulo
    ch.width, ch.height = w, h
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.legend.position = "b"
    return ch


def _serie(ws, col: int, r0: int, r1: int, titulo: str, cor: str) -> Series:
    s = Series(Reference(ws, min_col=col, min_row=r0, max_row=r1), title=titulo)
    s.graphicalProperties.solidFill = cor
    s.graphicalProperties.line.solidFill = cor
    return s


def write_xlsx(rows: list[dict], path: str, meta: dict | None = None) -> None:
    meta = meta or {}
    wb = Workbook()
    wb.calculation.fullCalcOnLoad = True

    tem_ragas = any(k.startswith("ragas_") and k != "ragas_unavailable_reason" for r in rows for k in r)
    presentes = {k for r in rows for k in r}
    colunas = [
        c for c in COLUNAS
        if (tem_ragas or not c[0].startswith("ragas_"))
        and (c[0] not in _COLUNAS_OPCIONAIS or c[0] in presentes)
    ]
    idx = {c[0]: i + 1 for i, c in enumerate(colunas)}  # chave -> nº da coluna em Resultados
    n = len(rows)
    ultima = n + 1  # linha final de dados em Resultados (cabeçalho na 1)

    def R(key: str) -> str:
        L = get_column_letter(idx[key])
        return f"Resultados!${L}$2:${L}${max(ultima, 2)}"

    ordem_cat: list[str] = []
    for r in rows:
        if r.get("category") not in ordem_cat:
            ordem_cat.append(r.get("category"))
    cat_com_exec = [c for c in ordem_cat if any(
        r.get("category") == c and r.get("topology") in ("star", "hierarchical") for r in rows)]
    cat_sem_exec = [c for c in ordem_cat if c not in cat_com_exec]

    # ---------------- Resultados (dados brutos) ----------------
    ws_res = wb.active
    ws_res.title = "Resultados"
    _cabecalho(ws_res, 1, [c[1] for c in colunas])
    for i, r in enumerate(rows, start=2):
        for j, (key, _, _, fmt) in enumerate(colunas, start=1):
            cell = ws_res.cell(row=i, column=j, value=_limpa(r.get(key)))
            cell.font = _F_BASE
            if fmt:
                cell.number_format = fmt
    _larguras(ws_res, [c[2] for c in colunas])
    ws_res.freeze_panes = "E2"
    if n:
        tab = Table(displayName="TabelaResultados", ref=f"A1:{get_column_letter(len(colunas))}{ultima}")
        tab.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
        ws_res.add_table(tab)

    # ---------------- Cenarios ----------------
    ws_cen = wb.create_sheet("Cenarios")
    _titulo(ws_cen, "Cenários da bateria", "Lista completa dos cenários selecionados; 'Executou topologias' conta as linhas de Resultados com topologia star/hierarchical.")
    _cabecalho(ws_cen, 4, ["Cenário", "Categoria", "Tam. (chars)", "Execuções de topologia", "Pergunta completa"])
    cenarios = meta.get("scenarios")
    if not cenarios:
        vistos: dict[str, dict] = {}
        for r in rows:
            vistos.setdefault(r["scenario"], {"id": r["scenario"], "category": r.get("category"),
                                              "question": r.get("question", ""), "len": r.get("question_len")})
        cenarios = list(vistos.values())
    for i, c in enumerate(cenarios, start=5):
        ws_cen.cell(row=i, column=1, value=c["id"])
        ws_cen.cell(row=i, column=2, value=c["category"])
        ws_cen.cell(row=i, column=3, value=c.get("len") or len(c.get("question", "")))
        ws_cen.cell(row=i, column=4, value=f'=COUNTIFS({R("scenario")},$A{i},{R("topology")},"star",{R("wall_clock_ms")},">=0")+COUNTIFS({R("scenario")},$A{i},{R("topology")},"hierarchical",{R("wall_clock_ms")},">=0")')
        ws_cen.cell(row=i, column=5, value=_limpa(c.get("question", "")))
    if cenarios:
        _estilo_corpo(ws_cen, 5, 4 + len(cenarios), 1, 5, {3: "#,##0", 4: "0"})
    _larguras(ws_cen, [38, 24, 12, 16, 120])
    ws_cen.freeze_panes = "A5"

    # ---------------- InterpDados (1 linha por modelo × cenário) ----------------
    ws_id = wb.create_sheet("InterpDados")
    _titulo(ws_id, "Interpretação de intenção — dados por tentativa", "Uma linha por modelo × cenário (deduplicado de Resultados).")
    chaves_i = [("model", "Modelo"), ("category", "Categoria"), ("scenario", "Cenário"),
                ("question_len", "Tam. pergunta"), ("intent_success", "Interpretação OK"),
                ("intent_missing", "Motivo da recusa"), ("intent_elapsed_ms", "Tempo (ms)"),
                ("intent_tokens", "Tokens"), ("health_params_inferred", "Tema inferido"),
                ("date_range_inferred", "Período inferido"), ("resolved_health_params_count", "Nº indicadores")]
    _cabecalho(ws_id, 4, [h for _, h in chaves_i])
    dedup: dict[tuple, dict] = {}
    for r in rows:
        if r.get("intent_reused"):
            continue
        dedup.setdefault((r.get("model"), r.get("scenario")), r)
    linhas_i = list(dedup.values())
    for i, r in enumerate(linhas_i, start=5):
        for j, (k, _) in enumerate(chaves_i, start=1):
            ws_id.cell(row=i, column=j, value=_limpa(r.get(k)))
    ni0, ni1 = 5, 4 + len(linhas_i)
    if linhas_i:
        _estilo_corpo(ws_id, ni0, ni1, 1, len(chaves_i), {4: "#,##0", 7: "#,##0", 8: "#,##0"})
    _larguras(ws_id, [18, 24, 38, 12, 14, 18, 12, 10, 12, 14, 12])
    ws_id.freeze_panes = "A5"

    def I(col_letter: str) -> str:
        return f"InterpDados!${col_letter}${ni0}:${col_letter}${max(ni1, ni0)}"

    # ---------------- Interpretacao (resumo por categoria) ----------------
    ws_in = wb.create_sheet("Interpretacao")
    _titulo(ws_in, "Interpretação de intenção por categoria",
            "Premissa: 'Esperado' é uma suposição do relatório (ver coluna B), não uma regra do sistema. 'Misto' não é pontuado.")
    heads = ["Categoria", "Esperado", "Tentativas", "Interpretadas com sucesso", "Taxa de sucesso",
             "Recusa: fora de escopo", "Recusa: conversa", "Falha técnica (LLM)", "Outras recusas",
             "Tema inferido", "Período inferido", "Tempo médio (ms)", "Tokens médios", "Aderência ao esperado"]
    _cabecalho(ws_in, 4, heads)
    for i, cat in enumerate(ordem_cat, start=5):
        ws_in.cell(row=i, column=1, value=cat)
        ws_in.cell(row=i, column=2, value=ESPERADO.get(cat, ESPERADO_PADRAO))
        ws_in.cell(row=i, column=3, value=f'=COUNTIFS({I("B")},$A{i})')
        ws_in.cell(row=i, column=4, value=f'=COUNTIFS({I("B")},$A{i},{I("E")},TRUE)')
        ws_in.cell(row=i, column=5, value=f'=IFERROR(D{i}/C{i},"")')
        ws_in.cell(row=i, column=6, value=f'=COUNTIFS({I("B")},$A{i},{I("F")},"fora_de_escopo")')
        ws_in.cell(row=i, column=7, value=f'=COUNTIFS({I("B")},$A{i},{I("F")},"conversa")')
        ws_in.cell(row=i, column=8, value=f'=COUNTIFS({I("B")},$A{i},{I("F")},"llm_indisponivel")')
        ws_in.cell(row=i, column=9, value=f"=C{i}-D{i}-F{i}-G{i}-H{i}")
        ws_in.cell(row=i, column=10, value=f'=COUNTIFS({I("B")},$A{i},{I("I")},TRUE)')
        ws_in.cell(row=i, column=11, value=f'=COUNTIFS({I("B")},$A{i},{I("J")},TRUE)')
        ws_in.cell(row=i, column=12, value=f'=IFERROR(AVERAGEIFS({I("G")},{I("B")},$A{i}),"")')
        ws_in.cell(row=i, column=13, value=f'=IFERROR(AVERAGEIFS({I("H")},{I("B")},$A{i}),"")')
        ws_in.cell(row=i, column=14, value=(
            f'=IF(C{i}=0,"",IF(B{i}="Analisar",D{i}/C{i},IF(B{i}="Recusar",(C{i}-D{i})/C{i},'
            f'IF(B{i}="Conversa",G{i}/C{i},"n/a"))))'))
    nc = len(ordem_cat)
    if nc:
        _estilo_corpo(ws_in, 5, 4 + nc, 1, len(heads),
                      {5: "0%", 12: "#,##0", 13: "#,##0", 14: "0%"})
        lt = 5 + nc
        ws_in.cell(row=lt, column=1, value="Total").font = _F_BOLD
        for c in (3, 4, 6, 7, 8, 9, 10, 11):
            L = get_column_letter(c)
            cell = ws_in.cell(row=lt, column=c, value=f"=SUM({L}5:{L}{4 + nc})")
            cell.font = _F_BOLD
        ws_in.cell(row=lt, column=5, value=f'=IFERROR(D{lt}/C{lt},"")').font = _F_BOLD
        ws_in.cell(row=lt, column=5).number_format = "0%"
        ch = _barras("Taxa de interpretação com sucesso por categoria", "Taxa", 24, 10)
        ch.add_data(Reference(ws_in, min_col=5, min_row=4, max_row=4 + nc), titles_from_data=True)
        ch.set_categories(Reference(ws_in, min_col=1, min_row=5, max_row=4 + nc))
        ch.series[0].graphicalProperties.solidFill = COR_NEUTRA
        ch.y_axis.number_format = "0%"
        ch.y_axis.scaling.min = 0
        ch.y_axis.scaling.max = 1
        ch.legend = None
        ws_in.add_chart(ch, f"A{lt + 3}")
        ch2 = _barras("Tempo médio de interpretação por categoria (ms)", "ms", 24, 10)
        ch2.add_data(Reference(ws_in, min_col=12, min_row=4, max_row=4 + nc), titles_from_data=True)
        ch2.set_categories(Reference(ws_in, min_col=1, min_row=5, max_row=4 + nc))
        ch2.series[0].graphicalProperties.solidFill = COR_ESTRELA
        ch2.legend = None
        ws_in.add_chart(ch2, f"H{lt + 3}")
    _larguras(ws_in, [26, 11, 11, 14, 11, 13, 12, 13, 11, 11, 11, 13, 12, 14])
    ws_in.freeze_panes = "B5"

    # ---------------- PorCategoria ----------------
    ws_pc = wb.create_sheet("PorCategoria")
    _titulo(ws_pc, "Estrela vs. hierárquica por categoria de cenário",
            "Médias sobre as execuções de topologia (cenários que passaram da interpretação). "
            + (f"Sem execução de topologia (só recusa/conversa): {', '.join(cat_sem_exec)}." if cat_sem_exec else ""))
    heads = ["Categoria", "Execuções estrela", "Execuções hierárquica",
             "Tempo médio estrela (ms)", "Tempo médio hierárquica (ms)", "Δ tempo (hier vs estrela)",
             "Tokens médios estrela", "Tokens médios hierárquica", "Δ tokens (hier vs estrela)",
             "Mensagens estrela", "Mensagens hierárquica", "Texto médio estrela (chars)", "Texto médio hierárquica (chars)"]
    _cabecalho(ws_pc, 4, heads)
    for i, cat in enumerate(cat_com_exec, start=5):
        ws_pc.cell(row=i, column=1, value=cat)
        for col, topo in ((2, "star"), (3, "hierarchical")):
            ws_pc.cell(row=i, column=col, value=f'=COUNTIFS({R("category")},$A{i},{R("topology")},"{topo}",{R("wall_clock_ms")},">=0")')

        def media(key: str, topo: str) -> str:
            return f'=IFERROR(AVERAGEIFS({R(key)},{R("category")},$A{i},{R("topology")},"{topo}"),"")'
        ws_pc.cell(row=i, column=4, value=media("wall_clock_ms", "star"))
        ws_pc.cell(row=i, column=5, value=media("wall_clock_ms", "hierarchical"))
        ws_pc.cell(row=i, column=6, value=f'=IF(OR(D{i}="",E{i}="",D{i}=0),"",E{i}/D{i}-1)')
        ws_pc.cell(row=i, column=7, value=media("total_tokens", "star"))
        ws_pc.cell(row=i, column=8, value=media("total_tokens", "hierarchical"))
        ws_pc.cell(row=i, column=9, value=f'=IF(OR(G{i}="",H{i}="",G{i}=0),"",H{i}/G{i}-1)')
        ws_pc.cell(row=i, column=10, value=media("message_count", "star"))
        ws_pc.cell(row=i, column=11, value=media("message_count", "hierarchical"))
        ws_pc.cell(row=i, column=12, value=media("texto_analise_len", "star"))
        ws_pc.cell(row=i, column=13, value=media("texto_analise_len", "hierarchical"))
    k = len(cat_com_exec)
    if k:
        _estilo_corpo(ws_pc, 5, 4 + k, 1, len(heads),
                      {4: "#,##0", 5: "#,##0", 6: "+0.0%;-0.0%;0.0%", 7: "#,##0", 8: "#,##0",
                       9: "+0.0%;-0.0%;0.0%", 10: "0.0", 11: "0.0", 12: "#,##0", 13: "#,##0"})
        cats = Reference(ws_pc, min_col=1, min_row=5, max_row=4 + k)
        for titulo, y, c_star, c_hier, anchor in (
            ("Tempo médio de parede por categoria", "ms", 4, 5, f"A{k + 8}"),
            ("Tokens médios por categoria", "tokens", 7, 8, f"H{k + 8}"),
        ):
            ch = _barras(titulo, y, 22, 10)
            ch.series.append(_serie(ws_pc, c_star, 5, 4 + k, "Estrela", COR_ESTRELA))
            ch.series.append(_serie(ws_pc, c_hier, 5, 4 + k, "Hierárquica", COR_HIER))
            ch.set_categories(cats)
            ws_pc.add_chart(ch, anchor)
    _larguras(ws_pc, [26, 11, 13, 14, 14, 14, 13, 14, 14, 12, 14, 14, 14])
    ws_pc.freeze_panes = "B5"

    # ---------------- EstrelaVsHier (pareado por modelo × cenário) ----------------
    ws_ev = wb.create_sheet("EstrelaVsHier")
    _titulo(ws_ev, "Comparação pareada: estrela vs. hierárquica",
            "Uma linha por modelo × cenário em que as duas topologias rodaram. Δ% = hierárquica / estrela − 1 (positivo = hierárquica maior).")
    heads = ["Modelo", "Categoria", "Cenário",
             "Tempo estrela (ms)", "Tempo hierárquica (ms)", "Δ tempo (ms)", "Δ tempo (%)", "Menor tempo",
             "Tokens estrela", "Tokens hierárquica", "Δ tokens", "Δ tokens (%)", "Menos tokens",
             "Mensagens estrela", "Mensagens hierárquica", "Texto estrela (chars)", "Texto hierárquica (chars)",
             "Completude estrela", "Completude hierárquica", "Fidelidade estrela", "Fidelidade hierárquica",
             "Correlações/anomalias idênticas"]
    _cabecalho(ws_ev, 4, heads)
    linha_de = {(r.get("model"), r.get("scenario"), r.get("topology")): r for r in rows}
    por_chave: dict[tuple, set] = {}
    for r in rows:
        if (r.get("topology") in ("star", "hierarchical") and r.get("wall_clock_ms") is not None
                and r.get("fault") in (None, "nenhuma")):
            por_chave.setdefault((r["model"], r["category"], r["scenario"]), set()).add(r["topology"])
    pares = [k for k, v in por_chave.items() if v == {"star", "hierarchical"}]
    e0, e1 = 5, 4 + len(pares)
    if pares:
        def soma(key: str, topo: str, i: int) -> str:
            return (f'=SUMIFS({R(key)},{R("model")},$A{i},{R("scenario")},$C{i},{R("topology")},"{topo}")')
        for i, (m, cat, sc) in enumerate(pares, start=5):
            ws_ev.cell(row=i, column=1, value=m)
            ws_ev.cell(row=i, column=2, value=cat)
            ws_ev.cell(row=i, column=3, value=sc)
            ws_ev.cell(row=i, column=4, value=soma("wall_clock_ms", "star", i))
            ws_ev.cell(row=i, column=5, value=soma("wall_clock_ms", "hierarchical", i))
            ws_ev.cell(row=i, column=6, value=f"=E{i}-D{i}")
            ws_ev.cell(row=i, column=7, value=f'=IF(D{i}=0,"",E{i}/D{i}-1)')
            ws_ev.cell(row=i, column=8, value=f'=IF(D{i}<E{i},"Estrela",IF(E{i}<D{i},"Hierárquica","Igual"))')
            ws_ev.cell(row=i, column=9, value=soma("total_tokens", "star", i))
            ws_ev.cell(row=i, column=10, value=soma("total_tokens", "hierarchical", i))
            ws_ev.cell(row=i, column=11, value=f"=J{i}-I{i}")
            ws_ev.cell(row=i, column=12, value=f'=IF(I{i}=0,"",J{i}/I{i}-1)')
            ws_ev.cell(row=i, column=13, value=f'=IF(I{i}<J{i},"Estrela",IF(J{i}<I{i},"Hierárquica","Igual"))')
            ws_ev.cell(row=i, column=14, value=soma("message_count", "star", i))
            ws_ev.cell(row=i, column=15, value=soma("message_count", "hierarchical", i))
            ws_ev.cell(row=i, column=16, value=soma("texto_analise_len", "star", i))
            ws_ev.cell(row=i, column=17, value=soma("texto_analise_len", "hierarchical", i))
            # Valores do payload de qualidade do sistema, copiados de Resultados
            # (vazio = não medido, diferente de zero).
            rs, rh = linha_de[(m, sc, "star")], linha_de[(m, sc, "hierarchical")]
            ws_ev.cell(row=i, column=18, value=rs.get("completude"))
            ws_ev.cell(row=i, column=19, value=rh.get("completude"))
            ws_ev.cell(row=i, column=20, value=rs.get("ragas_faithfulness"))
            ws_ev.cell(row=i, column=21, value=rh.get("ragas_faithfulness"))
            ident = rs.get("consistencia_identica")
            ws_ev.cell(row=i, column=22, value=None if ident is None else ("sim" if ident else "NÃO"))
        _estilo_corpo(ws_ev, e0, e1, 1, len(heads),
                      {4: "#,##0", 5: "#,##0", 6: "+#,##0;-#,##0;0", 7: "+0.0%;-0.0%;0.0%",
                       9: "#,##0", 10: "#,##0", 11: "+#,##0;-#,##0;0", 12: "+0.0%;-0.0%;0.0%",
                       14: "0", 15: "0", 16: "#,##0", 17: "#,##0", 18: "0.000", 19: "0.000", 20: "0.000", 21: "0.000"})
        w = max(30, min(len(pares) * 0.9, 80))
        for titulo, y, c_star, c_hier, anchor in (
            ("Tempo de parede por cenário", "ms", 4, 5, f"A{e1 + 3}"),
            ("Tokens totais por cenário", "tokens", 9, 10, f"A{e1 + 24}"),
        ):
            ch = _barras(titulo, y, w, 11)
            ch.series.append(_serie(ws_ev, c_star, e0, e1, "Estrela", COR_ESTRELA))
            ch.series.append(_serie(ws_ev, c_hier, e0, e1, "Hierárquica", COR_HIER))
            ch.set_categories(Reference(ws_ev, min_col=3, min_row=e0, max_row=e1))
            ws_ev.add_chart(ch, anchor)
    else:
        ws_ev["A5"] = "Nenhum cenário com as duas topologias executadas nesta rodada."
        ws_ev["A5"].font = _F_NOTE
    _larguras(ws_ev, [18, 24, 38, 13, 14, 12, 11, 13, 12, 14, 11, 11, 13, 12, 14, 13, 14, 11, 12, 11, 12, 14])
    ws_ev.freeze_panes = "D5"

    def E(col_letter: str) -> str:
        return f"EstrelaVsHier!${col_letter}${e0}:${col_letter}${max(e1, e0)}"

    # ---------------- Resumo (primeira aba) ----------------
    ws = wb.create_sheet("Resumo", 0)
    _titulo(ws, "Skopos — bateria de cenários: estrela vs. hierárquica")
    modelos = meta.get("models") or sorted({r.get("model") for r in rows if r.get("model")})
    ws["A2"] = (f"Gerado em {meta.get('generated_at') or datetime.now().strftime('%d/%m/%Y %H:%M')} · "
                f"modelos: {', '.join(map(str, modelos))} · {len(cenarios)} cenário(s) · {n} linha(s) em Resultados")
    ws["A2"].font = _F_NOTE

    ws["A4"] = "Médias gerais por topologia (sem falhas injetadas)"
    ws["A4"].font = _F_SUB
    _cabecalho(ws, 5, ["Métrica", "Estrela", "Hierárquica", "Δ (hier − estrela)", "Δ %"])
    metricas = [
        ("Execuções concluídas", None, "0"),
        ("Execuções com erro", "__erros__", "0"),
        ("Tempo de parede médio (ms)", "wall_clock_ms", "#,##0"),
        ("Tokens totais médios", "total_tokens", "#,##0"),
        ("Tokens de prompt médios", "prompt_tokens", "#,##0"),
        ("Tokens de resposta médios", "completion_tokens", "#,##0"),
        ("Chamadas LLM médias", "llm_call_count", "0.0"),
        ("Overhead de coordenação médio (%)", "overhead_percent", "0.00"),
        ("Invocações de agente médias", "agent_invocations", "0.0"),
        ("Mensagens médias", "message_count", "0.0"),
        ("Tamanho médio do texto (chars)", "texto_analise_len", "#,##0"),
        ("Completude média (0-1)", "completude", "0.000"),
        ("Cobertura de resultado R1 média (0-1)", "r1_score", "0.000"),
        ("Análises bem-sucedidas (outcome)", "sucesso", "0%"),
        ("Dentro do orçamento de tempo", "dentro_do_orcamento_de_tempo", "0%"),
    ]
    if tem_ragas:
        metricas += [("RAGAS faithfulness médio", "ragas_faithfulness", "0.000"),
                     ("RAGAS answer relevancy médio", "ragas_answer_relevancy", "0.000"),
                     ("RAGAS context relevance médio", "ragas_context_relevance", "0.000")]
    linha_metrica: dict[str, int] = {}
    for i, (nome, key, fmt) in enumerate(metricas, start=6):
        ws.cell(row=i, column=1, value=nome)
        for col, topo in ((2, "star"), (3, "hierarchical")):
            if key == "__erros__":
                f = (f'=COUNTIFS({R("topology")},"{topo}",{R("fault")},"nenhuma")'
                     f'-COUNTIFS({R("topology")},"{topo}",{R("fault")},"nenhuma",{R("wall_clock_ms")},">=0")')
            elif key is None:
                f = f'=COUNTIFS({R("topology")},"{topo}",{R("fault")},"nenhuma",{R("wall_clock_ms")},">=0")'
            else:
                f = f'=IFERROR(AVERAGEIFS({R(key)},{R("topology")},"{topo}",{R("fault")},"nenhuma"),"")'
            ws.cell(row=i, column=col, value=f)
        ws.cell(row=i, column=4, value=f'=IF(OR(B{i}="",C{i}=""),"",C{i}-B{i})')
        ws.cell(row=i, column=5, value=f'=IF(OR(B{i}="",C{i}="",B{i}=0),"",C{i}/B{i}-1)')
        linha_metrica[key or "n"] = i
    fm = len(metricas)
    _estilo_corpo(ws, 6, 5 + fm, 1, 5, {5: "+0.0%;-0.0%;0.0%"})
    for i, (_, _, fmt) in enumerate(metricas, start=6):
        for c in (2, 3, 4):
            ws.cell(row=i, column=c).number_format = fmt
    ws.cell(row=6, column=4).number_format = "+0;-0;0"

    r = 7 + fm
    ws.cell(row=r, column=1, value="Comparação nos cenários pareados").font = _F_SUB
    _cabecalho(ws, r + 1, ["Grandeza", "Menor na estrela", "Menor na hierárquica", "Iguais", "Cenários pareados"])
    for j, (nome, L) in enumerate((("Tempo de parede", "H"), ("Tokens totais", "M")), start=r + 2):
        ws.cell(row=j, column=1, value=nome)
        ws.cell(row=j, column=2, value=f'=COUNTIF({E(L)},"Estrela")')
        ws.cell(row=j, column=3, value=f'=COUNTIF({E(L)},"Hierárquica")')
        ws.cell(row=j, column=4, value=f'=COUNTIF({E(L)},"Igual")')
        ws.cell(row=j, column=5, value=f"=SUM(B{j}:D{j})")
    _estilo_corpo(ws, r + 2, r + 3, 1, 5, {2: "0", 3: "0", 4: "0", 5: "0"})

    r2 = r + 5
    ws.cell(row=r2, column=1, value="Interpretação de intenção (todas as tentativas)").font = _F_SUB
    _cabecalho(ws, r2 + 1, ["Indicador", "Valor"])
    itens = [
        ("Tentativas de interpretação", f"=Interpretacao!C{5 + len(ordem_cat)}", "0"),
        ("Interpretadas com sucesso", f"=Interpretacao!D{5 + len(ordem_cat)}", "0"),
        ("Taxa de sucesso", f"=Interpretacao!E{5 + len(ordem_cat)}", "0%"),
        ("Recusas por fora de escopo", f"=Interpretacao!F{5 + len(ordem_cat)}", "0"),
        ("Falhas técnicas do LLM", f"=Interpretacao!H{5 + len(ordem_cat)}", "0"),
        ("Tempo médio de interpretação (ms)", f'=IFERROR(AVERAGE({I("G")}),"")', "#,##0"),
    ]
    for j, (nome, f, fmt) in enumerate(itens, start=r2 + 2):
        ws.cell(row=j, column=1, value=nome)
        ws.cell(row=j, column=2, value=f).number_format = fmt
    _estilo_corpo(ws, r2 + 2, r2 + 1 + len(itens), 1, 2)
    for j, (_, _, fmt) in enumerate(itens, start=r2 + 2):
        ws.cell(row=j, column=2).number_format = fmt

    rn = r2 + len(itens) + 3
    notas = [
        "Como ler: Δ % positivo = a hierárquica gastou mais que a estrela. Tempo = parede da execução da topologia (exclui a interpretação de intenção).",
        "Tokens vêm do TokenBucket de cada topologia (todas as chamadas LLM dos agentes). Interpretação de intenção é contabilizada à parte (aba Interpretacao).",
        "Overhead de coordenação e mensagens vêm de core.quality_metrics (compute_coordination_overhead / compute_communication_volume).",
        "Linhas com topologia '-' são cenários recusados/conversa na interpretação: não rodaram nenhuma topologia.",
        "Abas: Resultados (dados brutos) · PorCategoria · EstrelaVsHier (pareado) · Interpretacao / InterpDados · Cenarios.",
    ]
    for j, t in enumerate(notas, start=rn):
        ws.cell(row=j, column=1, value=t).font = _F_NOTE

    for titulo, key, y, anchor in (
        ("Tempo de parede médio (ms)", "wall_clock_ms", "ms", "G4"),
        ("Tokens totais médios", "total_tokens", "tokens", "G22"),
        ("Mensagens médias", "message_count", "mensagens", "N4"),
        ("Tamanho médio do texto (chars)", "texto_analise_len", "chars", "N22"),
    ):
        lr = linha_metrica.get(key)
        if lr is None:
            continue
        ch = _barras(titulo, y, 11.5, 8)
        ch.series.append(_serie(ws, 2, lr, lr, "Estrela", COR_ESTRELA))
        ch.series.append(_serie(ws, 3, lr, lr, "Hierárquica", COR_HIER))
        ws.add_chart(ch, anchor)
    _larguras(ws, [40, 14, 16, 18, 12])

    wb.save(path)


def _carregar_jsonl(path: str) -> list[dict]:
    with open(path, encoding="utf-8") as f:
        return [json.loads(l) for l in f if l.strip()]


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print("uso: python -m scripts.batch_report entrada.jsonl saida.xlsx")
        raise SystemExit(2)
    dados = _carregar_jsonl(sys.argv[1])
    write_xlsx(dados, sys.argv[2])
    print(f"{len(dados)} linhas -> {sys.argv[2]}")
