"""
Consolida N rodadas da bateria (`scripts.batch_eval`) e produz as estatísticas,
tabelas e figuras usadas no capítulo de resultados.

Cada rodada é um `.jsonl` gravado por `batch_eval --out` (uma linha por
cenário × topologia) e, ao lado, o `.payloads.jsonl` correspondente. Nada é
recalculado a partir do sistema: o script só agrega o que as rodadas mediram.

O que sai em `--out`:

  analise.json   todos os números (o texto do capítulo cita este arquivo)
  analise.xlsx   as mesmas tabelas, mais a planilha por pergunta (valores das duas topologias)
  fig_*.png      figuras

Tratamento estatístico: média, mediana e quartis por topologia; contagem das
perguntas em que a hierárquica ficou acima, igual ou abaixo da estrela; e o
teste de postos sinalizados de Wilcoxon sobre as diferenças pareadas por
pergunta (com mais de uma execução, sobre as medianas de cada pergunta). Sem
correção para comparações múltiplas; todos os p são gravados.

Uso (na raiz do repositório; requer numpy, scipy, matplotlib e openpyxl):

    python backend/scripts/analise_rodadas.py \
        --rodada resultados_ui_completo.jsonl \
        --rodada resultados_rodada2.jsonl --rodada resultados_rodada3.jsonl \
        --rodada resultados_rodada4.jsonl --rodada resultados_rodada5.jsonl \
        --substituir-na-primeira rerun_afetadas.jsonl \
        --out analise_rodadas
"""

from __future__ import annotations

import argparse
import json
import os
import statistics as st
from collections import Counter, defaultdict
from itertools import combinations

import numpy as np
from scipy import stats

TOPOS = ("star", "hierarchical")
NOME_TOPO = {"star": "Estrela", "hierarchical": "Hierárquica"}
# Limiar de indiferença: diferença de fidelidade abaixo da qual o próprio
# avaliador não distingue os textos. Serve de régua para a análise do ruído.
LIMIAR_INDIFERENCA = 0.05

# (chave no jsonl, rótulo, casas decimais)
METRICAS = [
    ("wall_clock_ms", "Tempo de parede (ms)", 0),
    ("tempo_total_ms", "Tempo até o texto (ms)", 0),
    ("overhead_ms", "Custo de coordenação (ms)", 2),
    ("overhead_percent", "Custo de coordenação (%)", 2),
    ("overhead_deliberacao_ms", "Coordenação: deliberação LLM (ms)", 0),
    ("overhead_estrutural_ms", "Coordenação: estrutura (ms)", 2),
    ("total_tokens", "Tokens totais", 0),
    ("prompt_tokens", "Tokens de entrada", 0),
    ("completion_tokens", "Tokens de saída", 0),
    ("llm_call_count", "Chamadas ao modelo", 0),
    ("agent_invocations", "Agentes acionados", 0),
    ("message_count", "Mensagens", 0),
    ("mensagens_texto", "Mensagens de texto", 0),
    ("ferramentas_acionadas", "Ferramentas acionadas", 0),
    ("payload_records", "Registros transportados", 0),
    ("texto_analise_len", "Tamanho do laudo (caracteres)", 0),
    ("completude", "Completude", 3),
    ("ragas_faithfulness", "Fidelidade", 3),
    ("ragas_answer_relevancy", "Relevância da resposta", 3),
    ("ragas_context_relevance", "Relevância do contexto", 3),
    ("fidelidade_numerica", "Fidelidade numérica", 3),
]
METRICAS_TESTE = [
    "wall_clock_ms", "tempo_total_ms", "total_tokens", "completion_tokens", "message_count",
    "texto_analise_len", "completude", "ragas_faithfulness", "ragas_answer_relevancy",
    "fidelidade_numerica",
]


# --- leitura ---------------------------------------------------------------

def carregar(path: str) -> list[dict]:
    with open(path, encoding="utf-8") as f:
        return [json.loads(linha) for linha in f if linha.strip()]


def carregar_payloads(path_jsonl: str) -> dict[str, dict]:
    path = os.path.splitext(path_jsonl)[0] + ".payloads.jsonl"
    if not os.path.exists(path):
        return {}
    return {d["scenario"]: d for d in carregar(path)}


class Rodada:
    def __init__(self, numero: int, path: str):
        self.numero = numero
        self.path = path
        self.rows = carregar(path)
        self.payloads = carregar_payloads(path)
        self.substituidos: list[str] = []

    def substituir(self, path: str) -> None:
        """Troca os cenários presentes em `path` pelas linhas desse arquivo."""
        novas = carregar(path)
        cenarios = {r["scenario"] for r in novas}
        ordem = {}
        for r in self.rows:
            ordem.setdefault(r["scenario"], len(ordem))
        self.rows = [r for r in self.rows if r["scenario"] not in cenarios] + novas
        self.rows.sort(key=lambda r: ordem.get(r["scenario"], len(ordem)))  # estável: mantém estrela antes de hierárquica
        self.payloads.update(carregar_payloads(path))
        self.substituidos = sorted(cenarios)

    def tentativa(self, cenario: str) -> dict | None:
        return next((r for r in self.rows if r["scenario"] == cenario), None)

    def par(self, cenario: str) -> dict[str, dict] | None:
        """As duas linhas (estrela, hierárquica) de uma análise com métricas."""
        linhas = {r["topology"]: r for r in self.rows
                  if r["scenario"] == cenario and r.get("topology") in TOPOS}
        if len(linhas) != 2:
            return None
        if any(r.get("run_error") or r.get("wall_clock_ms") is None for r in linhas.values()):
            return None
        return linhas


# --- estatística -----------------------------------------------------------

def quartis(xs: list[float]) -> dict:
    xs = sorted(x for x in xs if x is not None)
    if not xs:
        return {"n": 0}
    q1, q3 = (np.percentile(xs, 25), np.percentile(xs, 75))
    return {"n": len(xs), "mediana": float(st.median(xs)), "q1": float(q1), "q3": float(q3),
            "min": float(xs[0]), "max": float(xs[-1]), "media": float(st.mean(xs))}


def wilcoxon_pareado(difs: list[float]) -> float | None:
    difs = [d for d in difs if d is not None]
    if len(difs) < 6 or all(d == 0 for d in difs):
        return None
    return float(stats.wilcoxon(difs, zero_method="wilcox").pvalue)


def percentis(xs: list[float], ps=(50, 75, 90, 95)) -> dict:
    return {f"p{p}": float(np.percentile(xs, p)) for p in ps} if xs else {}


# --- análise ---------------------------------------------------------------

def analisar(rodadas: list[Rodada]) -> dict:
    R = len(rodadas)
    base = rodadas[0]
    cenarios = []  # perguntas distintas (sem as repetições com falha injetada)
    for r in base.rows:
        if r["category"] != "falha_injetada" and r["scenario"] not in cenarios:
            cenarios.append(r["scenario"])
    categoria = {c: base.tentativa(c)["category"] for c in cenarios}
    out: dict = {"rodadas": R, "arquivos": [rd.path for rd in rodadas],
                 "substituidos_na_rodada_1": base.substituidos, "perguntas": len(cenarios)}

    # ---- admissão ----
    adm = {}
    for c in cenarios:
        status = []
        for rd in rodadas:
            t = rd.tentativa(c)
            status.append("ausente" if t is None else ("admitida" if t.get("intent_success") else t.get("intent_missing") or "recusada"))
        adm[c] = status
    n_adm = [sum(1 for s in v if s == "admitida") for v in adm.values()]
    motivos = Counter(s for v in adm.values() for s in v if s != "admitida")
    instaveis = {c: v for c, v in adm.items() if 0 < sum(1 for s in v if s == "admitida") < R}
    tempos_i, tokens_i = [], []
    for rd in rodadas:
        for c in cenarios:
            t = rd.tentativa(c)
            if t and t.get("intent_elapsed_ms") is not None and not t.get("intent_reused"):
                tempos_i.append(t["intent_elapsed_ms"])
                tokens_i.append(t["intent_tokens"])
    resolucao = {}
    for c in cenarios:
        resolucao[c] = []
        for rd in rodadas:
            t = rd.tentativa(c)
            resolucao[c].append([t.get("resolved_date_from"), t.get("resolved_date_to"), t.get("resolved_health_params_count")]
                                if t and t.get("intent_success") else None)
    out["admissao"] = {
        "por_pergunta": adm,
        "resolucao_por_pergunta": resolucao,
        "categoria": categoria,
        "tentativas": len(cenarios) * R,
        "admitidas": sum(n_adm),
        "admitidas_por_rodada": [sum(1 for v in adm.values() if v[i] == "admitida") for i in range(R)],
        "sempre_admitidas": sum(1 for n in n_adm if n == R),
        "nunca_admitidas": sum(1 for n in n_adm if n == 0),
        "instaveis": instaveis,
        "motivos_de_recusa": dict(motivos),
        "tempo_ms": quartis(tempos_i),
        "tokens": quartis(tokens_i),
    }

    # ---- conjunto pareado ----
    pares = {c: [rd.par(c) for rd in rodadas] for c in cenarios}
    Q = [c for c in cenarios if all(p is not None for p in pares[c])]
    parciais = {c: sum(1 for p in pares[c] if p is not None) for c in cenarios
                if any(p is not None for p in pares[c]) and c not in Q}
    out["pareado"] = {"perguntas_em_todas_as_rodadas": len(Q), "analises": len(Q) * R,
                      "perguntas_em_parte_das_rodadas": parciais,
                      "analises_totais_com_metricas": sum(1 for c in cenarios for p in pares[c] if p)}

    def vals(c: str, topo: str, m: str) -> list:
        return [p[topo].get(m) for p in pares[c]]

    # ---- estabilidade da interpretação ----
    resol = {}
    for c in Q:
        chaves = [(p["star"].get("resolved_date_from"), p["star"].get("resolved_date_to"),
                   tuple(sorted((p["star"].get("resolved_health_params") or "").split(","))))
                  for p in pares[c]]
        resol[c] = chaves
    estaveis = [c for c in Q if len(set(resol[c])) == 1]
    periodo_var = [c for c in Q if len({(k[0], k[1]) for k in resol[c]}) > 1]
    tema_var = [c for c in Q if len({k[2] for k in resol[c]}) > 1]
    out["interpretacao"] = {
        "resolucao_identica_em_todas": len(estaveis),
        "periodo_variou": periodo_var,
        "indicadores_variaram": tema_var,
        "tema_inferido": sum(1 for c in Q if any(p["star"].get("health_params_inferred") for p in pares[c])),
        "periodo_inferido": sum(1 for c in Q if any(p["star"].get("date_range_inferred") for p in pares[c])),
    }

    # ---- descritiva e testes por métrica ----
    metricas = {}
    por_pergunta = {c: {"categoria": categoria[c]} for c in Q}
    for chave, rotulo, _ in METRICAS:
        agrup = {t: [v for c in Q for v in vals(c, t, chave) if v is not None] for t in TOPOS}
        med = {t: {c: (st.median([v for v in vals(c, t, chave) if v is not None])
                       if any(v is not None for v in vals(c, t, chave)) else None) for c in Q} for t in TOPOS}
        difs = [med["hierarchical"][c] - med["star"][c] for c in Q
                if med["star"][c] is not None and med["hierarchical"][c] is not None]
        # diferença dentro da mesma análise (mesma rodada, mesma pergunta)
        dif_analise = [p["hierarchical"][chave] - p["star"][chave] for c in Q for p in pares[c]
                       if p["star"].get(chave) is not None and p["hierarchical"].get(chave) is not None]
        m = {
            "rotulo": rotulo,
            "star": quartis(agrup["star"]), "hierarchical": quartis(agrup["hierarchical"]),
            "mediana_das_medianas": {
                t: (float(st.median(vs)) if (vs := [v for v in med[t].values() if v is not None]) else None)
                for t in TOPOS
            },
            "perguntas_hier_maior": sum(1 for d in difs if d > 0),
            "perguntas_hier_menor": sum(1 for d in difs if d < 0),
            "perguntas_iguais": sum(1 for d in difs if d == 0),
            "wilcoxon_p": wilcoxon_pareado(difs),
            "dif_mediana_por_pergunta": float(st.median(difs)) if difs else None,
            "analises_hier_maior": sum(1 for d in dif_analise if d > 0),
            "analises_hier_menor": sum(1 for d in dif_analise if d < 0),
            "analises_iguais": sum(1 for d in dif_analise if d == 0),
            "dif_mediana_por_analise": float(st.median(dif_analise)) if dif_analise else None,
        }
        if chave in METRICAS_TESTE:
            for c in Q:
                por_pergunta[c][chave] = {"star": med["star"][c], "hierarchical": med["hierarchical"][c]}
        metricas[chave] = m
    out["metricas"] = metricas
    out["por_pergunta"] = por_pergunta

    # ---- invariantes dentro da análise ----
    analises = [(c, i, pares[c][i]) for c in Q for i in range(R)]
    iguais = lambda k: sum(1 for _, _, p in analises if p["star"].get(k) == p["hierarchical"].get(k))
    delta_inv = Counter(p["hierarchical"]["agent_invocations"] - p["star"]["agent_invocations"] for _, _, p in analises)
    # análises em que as topologias não partiram dos mesmos dados (o planejamento
    # de consulta dos agentes de domínio é uma chamada ao modelo feita em cada topologia)
    divergentes = [
        {"pergunta": c, "rodada": rodadas[i].numero,
         "consistencia_identica": p["star"].get("consistencia_identica"),
         **{f"{k}|{t}": p[t].get(k) for k in ("payload_records", "llm_call_count", "prompt_tokens", "r1_completed") for t in TOPOS}}
        for c, i, p in analises
        if p["star"].get("consistencia_identica") != 1
        or p["star"]["payload_records"] != p["hierarchical"]["payload_records"]
        or p["star"]["llm_call_count"] != p["hierarchical"]["llm_call_count"]
    ]
    chaves_div = {(d["pergunta"], d["rodada"]) for d in divergentes}
    convergentes = [(c, i, p) for c, i, p in analises if (c, rodadas[i].numero) not in chaves_div]
    out["invariantes"] = {
        "analises": len(analises),
        "divergentes": divergentes,
        "convergentes": len(convergentes),
        "tokens_de_entrada_iguais_nas_convergentes": sum(
            1 for _, _, p in convergentes if p["star"]["prompt_tokens"] == p["hierarchical"]["prompt_tokens"]),
        "dif_max_tokens_de_entrada_nas_convergentes": max(
            (abs(p["star"]["prompt_tokens"] - p["hierarchical"]["prompt_tokens"]) for _, _, p in convergentes), default=0),
        "analises_com_fidelidade_ausente": sum(
            1 for _, _, p in analises if p["star"].get("ragas_faithfulness") is None or p["hierarchical"].get("ragas_faithfulness") is None),
        "consistencia_deterministica": sum(1 for _, _, p in analises if p["star"].get("consistencia_identica") == 1),
        "tokens_de_entrada_iguais": iguais("prompt_tokens"),
        "chamadas_iguais": iguais("llm_call_count"),
        "registros_iguais": iguais("payload_records"),
        "delta_agentes_acionados": {str(k): v for k, v in sorted(delta_inv.items())},
        "dentro_do_orcamento_de_tempo": {t: sum(p[t]["dentro_do_orcamento_de_tempo"] for _, _, p in analises) for t in TOPOS},
        "sucesso": {t: sum(p[t]["sucesso"] for _, _, p in analises) for t in TOPOS},
        "r1_completo": {t: sum(1 for _, _, p in analises if p[t]["r1_completed"] == 7) for t in TOPOS},
        "r1_igual_entre_topologias": iguais("r1_completed"),
        "eventos_de_erro": {t: sum(p[t]["error_events"] for _, _, p in analises) for t in TOPOS},
        "fidelidade_nao_medida": {t: sum(1 for _, _, p in analises if p[t].get("ragas_faithfulness") is None) for t in TOPOS},
    }
    out["invariantes"]["dif_max_tokens_de_entrada"] = max(
        abs(p["star"]["prompt_tokens"] - p["hierarchical"]["prompt_tokens"]) for _, _, p in analises)
    out["invariantes"]["fracao_do_tempo_na_fase_de_dominio"] = {
        t: quartis([p[t]["fase_dominio_ms"] / p[t]["wall_clock_ms"] for _, _, p in analises if p[t]["wall_clock_ms"] > 0])
        for t in TOPOS}
    out["invariantes"]["hier_mais_rapida"] = sum(
        1 for _, _, p in analises if p["hierarchical"]["wall_clock_ms"] < p["star"]["wall_clock_ms"])
    # estabilidade entre rodadas dos dados recuperados (mesma pergunta, mesma resolução)
    out["invariantes"]["registros_iguais_entre_rodadas"] = sum(
        1 for c in estaveis if len({p["star"]["payload_records"] for p in pares[c]}) == 1)
    out["invariantes"]["perguntas_de_resolucao_estavel"] = len(estaveis)
    r1_incompleto = Counter()
    for c, _, p in analises:
        if p["star"]["r1_completed"] < 7:
            r1_incompleto[c] += 1
    out["invariantes"]["perguntas_com_r1_incompleto"] = dict(r1_incompleto)

    # ---- payloads: achados e contexto entregue ao avaliador ----
    n_corr, n_anom, ctx_iguais, ctx_total, cr_dif_com_ctx_igual, juiz = [], [], 0, 0, 0, Counter()
    ragas_erros = 0
    for rd in rodadas:
        for c in Q:
            pl = rd.payloads.get(c)
            if not pl:
                continue
            dc = pl["quality"]["quality"]["deterministic_consistency"]
            n_corr.append(dc.get("star_correlacoes_count", 0))
            n_anom.append(dc.get("star_anomalias_count", 0))
            rg = pl.get("ragas") or {}
            s, h = rg.get("star") or {}, rg.get("hierarchical") or {}
            if s.get("sample") and h.get("sample"):
                ctx_total += 1
                mesmo = s["sample"].get("n_contexts_total") == h["sample"].get("n_contexts_total")
                ctx_iguais += mesmo
                cs = ((s.get("metrics") or {}).get("context_relevance") or {}).get("score")
                ch = ((h.get("metrics") or {}).get("context_relevance") or {}).get("score")
                if mesmo and cs is not None and ch is not None and cs != ch:
                    cr_dif_com_ctx_igual += 1
            for lado in (s, h):
                if lado.get("judge"):
                    juiz[(lado["judge"].get("model"), lado["judge"].get("embedding_model_used"), lado.get("version"))] += 1
                ragas_erros += len(lado.get("errors") or [])
    # completude: a métrica vale 1 por convenção quando não há achados a cobrir
    comp_det = {}
    for t, sufixo in (("star", "star"), ("hierarchical", "hier")):
        todos, com_achados, sem_achados = [], [], 0
        for c in Q:
            for i, rd in enumerate(rodadas):
                v = pares[c][i][t].get("completude")
                pl = rd.payloads.get(c)
                if v is None or not pl:
                    continue
                dc = pl["quality"]["quality"]["deterministic_consistency"]
                n_achados = dc.get(f"{sufixo}_correlacoes_count", 0) + dc.get(f"{sufixo}_anomalias_count", 0)
                todos.append(v)
                if n_achados > 0:
                    com_achados.append(v)
                else:
                    sem_achados += 1
        comp_det[t] = {
            "analises": len(todos), "total": sum(1 for v in todos if v == 1),
            "sem_achados": sem_achados, "com_achados": quartis(com_achados),
            "total_com_achados": sum(1 for v in com_achados if v == 1),
        }
    out["completude_detalhe"] = comp_det
    out["payloads"] = {
        "correlacoes_por_analise": quartis(n_corr), "anomalias_por_analise": quartis(n_anom),
        "analises_com_amostra_ragas": ctx_total, "mesmo_numero_de_contextos": ctx_iguais,
        "relevancia_do_contexto_diferente_com_contexto_igual": cr_dif_com_ctx_igual,
        "avaliador": [{"modelo": k[0], "embeddings": k[1], "ragas": k[2], "avaliacoes": v} for k, v in juiz.items()],
        "erros_ragas": ragas_erros,
    }

    # ---- ruído da fidelidade: entre topologias × entre repetições ----
    ruido = {}
    for chave in ("ragas_faithfulness", "ragas_answer_relevancy", "completude"):
        entre_topo = [abs(p["hierarchical"][chave] - p["star"][chave]) for _, _, p in analises
                      if p["star"].get(chave) is not None and p["hierarchical"].get(chave) is not None]
        entre_rep, entre_rep_est = [], []
        for c in Q:
            for t in TOPOS:
                v = [x for x in vals(c, t, chave) if x is not None]
                d = [abs(a - b) for a, b in combinations(v, 2)]
                entre_rep += d
                if c in estaveis:
                    entre_rep_est += d
        ruido[chave] = {
            "entre_topologias": {**quartis(entre_topo), **percentis(entre_topo),
                                 "acima_do_limiar_de_indiferenca": sum(1 for d in entre_topo if d >= LIMIAR_INDIFERENCA)},
            "entre_repeticoes": {**quartis(entre_rep), **percentis(entre_rep),
                                 "acima_do_limiar_de_indiferenca": sum(1 for d in entre_rep if d >= LIMIAR_INDIFERENCA)},
            "entre_repeticoes_resolucao_estavel": {**quartis(entre_rep_est), **percentis(entre_rep_est)},
            "_entre_topo": entre_topo, "_entre_rep": entre_rep,
        }
    out["ruido"] = ruido
    zeros = Counter()
    for c, _, p in analises:
        for t in TOPOS:
            if p[t].get("ragas_answer_relevancy") == 0:
                zeros[(c, t)] += 1
    out["relevancia_zerada"] = {
        "total": sum(zeros.values()),
        "por_topologia": {t: sum(v for (c, tt), v in zeros.items() if tt == t) for t in TOPOS},
        "perguntas": sorted({c for c, _ in zeros}),
        "detalhe": {f"{c}|{t}": v for (c, t), v in sorted(zeros.items())},
        "mediana_sem_zeros": {t: float(st.median([p[t]["ragas_answer_relevancy"] for _, _, p in analises
                                                  if p[t].get("ragas_answer_relevancy")])) for t in TOPOS},
    }
    cr_iguais = sum(1 for _, _, p in analises
                    if p["star"].get("ragas_context_relevance") == p["hierarchical"].get("ragas_context_relevance"))
    out["relevancia_do_contexto"] = {"iguais_entre_topologias": cr_iguais, "analises": len(analises)}

    # ---- amplitude temporal × amplitude de escopo ----
    def moda(xs):
        return Counter(xs).most_common(1)[0][0]
    n_ind = {c: moda([p["star"]["resolved_health_params_count"] for p in pares[c]]) for c in Q}
    anos = {c: moda([p["star"]["resolved_date_to"] - p["star"]["resolved_date_from"] + 1 for p in pares[c]]) for c in Q}
    faixa_t = lambda a: "curta (2 a 3 anos)" if a <= 3 else ("média (4 a 6 anos)" if a <= 6 else "longa (7 a 10 anos)")
    faixa_e = lambda n: "focal (1 indicador)" if n == 1 else ("intermediário (2 a 4)" if n <= 4 else "amplo (todos os sistemas)")
    def medq(c, t, m):
        vs = [v for v in vals(c, t, m) if v is not None]
        return st.median(vs) if vs else None
    celulas = defaultdict(list)
    for c in Q:
        celulas[(faixa_t(anos[c]), faixa_e(n_ind[c]))].append(c)
    estratos = []
    ordem_t = ["curta (2 a 3 anos)", "média (4 a 6 anos)", "longa (7 a 10 anos)"]
    ordem_e = ["focal (1 indicador)", "intermediário (2 a 4)", "amplo (todos os sistemas)"]
    for (ft, fe), cs in sorted(celulas.items(), key=lambda kv: (ordem_t.index(kv[0][0]), ordem_e.index(kv[0][1]))):
        linha = {"periodo": ft, "escopo": fe, "perguntas": len(cs)}
        for m in ("wall_clock_ms", "total_tokens", "message_count", "ragas_faithfulness", "completude"):
            for t in TOPOS:
                linha[f"{m}|{t}"] = quartis([medq(c, t, m) for c in cs])
        estratos.append(linha)
    por_escopo = []
    for fe in ("focal (1 indicador)", "intermediário (2 a 4)", "amplo (todos os sistemas)"):
        cs = [c for c in Q if faixa_e(n_ind[c]) == fe]
        linha = {"escopo": fe, "perguntas": len(cs)}
        for m in ("wall_clock_ms", "total_tokens", "llm_call_count", "message_count", "payload_records",
                  "ragas_faithfulness", "completude"):
            for t in TOPOS:
                linha[f"{m}|{t}"] = quartis([medq(c, t, m) for c in cs]) if cs else {"n": 0}
        por_escopo.append(linha)
    corr = {}
    for nome, x in (("n_indicadores", n_ind), ("anos", anos)):
        for m in ("wall_clock_ms", "total_tokens", "llm_call_count", "message_count", "payload_records",
                  "ragas_faithfulness", "completude"):
            for t in TOPOS:
                xy = [(x[c], y) for c in Q if (y := medq(c, t, m)) is not None]
                r = stats.spearmanr([a for a, _ in xy], [b for _, b in xy])
                corr[f"{nome}|{m}|{t}"] = {"rho": float(r.statistic), "p": float(r.pvalue)}
    out["escopo"] = {
        "distribuicao_indicadores": {str(k): v for k, v in sorted(Counter(n_ind.values()).items())},
        "distribuicao_anos": {str(k): v for k, v in sorted(Counter(anos.values()).items())},
        "celulas": estratos, "por_escopo": por_escopo, "spearman": corr,
        "_n_ind": n_ind, "_anos": anos,
    }

    # ---- por categoria ----
    cats = defaultdict(list)
    for c in Q:
        cats[categoria[c]].append(c)
    out["por_categoria"] = []
    for cat, cs in cats.items():
        linha = {"categoria": cat, "perguntas": len(cs), "analises": len(cs) * R}
        for m in ("wall_clock_ms", "total_tokens", "ragas_faithfulness", "ragas_answer_relevancy", "completude"):
            for t in TOPOS:
                linha[f"{m}|{t}"] = float(st.median([v for c in cs for v in vals(c, t, m) if v is not None]))
        out["por_categoria"].append(linha)

    # ---- por rodada ----
    out["por_rodada"] = []
    for i, rd in enumerate(rodadas):
        ps = [pares[c][i] for c in Q]
        linha = {"rodada": rd.numero, "arquivo": os.path.basename(rd.path),
                 "admitidas": out["admissao"]["admitidas_por_rodada"][i]}
        for m in ("wall_clock_ms", "total_tokens", "ragas_faithfulness", "ragas_answer_relevancy", "completude"):
            for t in TOPOS:
                linha[f"{m}|{t}"] = float(st.median([p[t][m] for p in ps if p[t].get(m) is not None]))
        topo_rows = [r for r in rd.rows if r.get("topology") in TOPOS]
        linha["tokens_topologias"] = sum(r.get("total_tokens") or 0 for r in topo_rows)
        vistos, tok_int = set(), 0
        for r in rd.rows:
            if r["scenario"] not in vistos and not r.get("intent_reused"):
                tok_int += r.get("intent_tokens") or 0
            vistos.add(r["scenario"])
        linha["tokens_interpretacao"] = tok_int
        linha["analises_executadas"] = len(topo_rows) // 2
        for sufixo in ("inicio", "fim"):
            p = os.path.splitext(rd.path)[0] + f".{sufixo}.txt"
            linha[sufixo] = open(p, encoding="utf-8").read().strip() if os.path.exists(p) else None
        out["por_rodada"].append(linha)
    out["campanha"] = {
        "analises_executadas": sum(l["analises_executadas"] for l in out["por_rodada"]),
        "tokens_topologias": sum(l["tokens_topologias"] for l in out["por_rodada"]),
        "tokens_interpretacao": sum(l["tokens_interpretacao"] for l in out["por_rodada"]),
    }

    return out


# --- saída -----------------------------------------------------------------

def _limpar(o):
    """Remove as chaves internas (prefixo `_`) e converte tipos numpy."""
    if isinstance(o, dict):
        return {str(k): _limpar(v) for k, v in o.items() if not str(k).startswith("_")}
    if isinstance(o, (list, tuple)):
        return [_limpar(v) for v in o]
    if isinstance(o, (np.floating, np.integer)):
        return o.item()
    return o


def gravar_xlsx(a: dict, path: str) -> None:
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumo"
    ws.append(["Métrica", "Estrela: mediana", "Estrela: Q1", "Estrela: Q3", "Hierárquica: mediana",
               "Hierárquica: Q1", "Hierárquica: Q3", "Perguntas hier > estrela", "Perguntas hier < estrela",
               "Perguntas iguais", "Wilcoxon p"])
    for chave, rotulo, _ in METRICAS:
        m = a["metricas"][chave]
        ws.append([rotulo, m["star"].get("mediana"), m["star"].get("q1"), m["star"].get("q3"),
                   m["hierarchical"].get("mediana"), m["hierarchical"].get("q1"), m["hierarchical"].get("q3"),
                   m["perguntas_hier_maior"], m["perguntas_hier_menor"], m["perguntas_iguais"], m["wilcoxon_p"]])
    ws = wb.create_sheet("PorPergunta")
    cab = ["Pergunta", "Categoria"]
    for chave in METRICAS_TESTE:
        rot = next(r for k, r, _ in METRICAS if k == chave)
        cab += [f"{rot}: estrela", f"{rot}: hierárquica"]
    ws.append(cab)
    for c, d in a["por_pergunta"].items():
        linha = [c, d["categoria"]]
        for chave in METRICAS_TESTE:
            x = d.get(chave) or {}
            linha += [x.get("star"), x.get("hierarchical")]
        ws.append(linha)
    ws = wb.create_sheet("Admissao")
    ws.append(["Pergunta", "Categoria"] + [f"Rodada {i + 1}" for i in range(a["rodadas"])])
    for c, v in a["admissao"]["por_pergunta"].items():
        ws.append([c, a["admissao"]["categoria"][c]] + v)
    ws = wb.create_sheet("PorRodada")
    chaves = list(a["por_rodada"][0])
    ws.append(chaves)
    for l in a["por_rodada"]:
        ws.append([l[k] for k in chaves])
    wb.save(path)


def gravar_figuras(a: dict, rodadas: list[Rodada], pasta: str) -> None:
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    azul, laranja = "#2a78d6", "#eb6834"  # estrela, hierárquica (paleta validada para daltonismo)
    tinta, cinza, grade = "#0b0b0b", "#52514e", "#e4e3df"
    plt.rcParams.update({
        "font.family": "Arial", "font.size": 9, "axes.edgecolor": cinza, "axes.labelcolor": tinta,
        "xtick.color": cinza, "ytick.color": cinza, "axes.spines.top": False, "axes.spines.right": False,
        "axes.grid": True, "grid.color": grade, "grid.linewidth": 0.6, "axes.axisbelow": True,
        "legend.frameon": False, "figure.dpi": 100, "savefig.dpi": 300, "savefig.bbox": "tight",
    })
    Q = list(a["por_pergunta"])
    n_ind = a["escopo"]["_n_ind"]
    pares = {c: [rd.par(c) for rd in rodadas] for c in Q}
    virgula = matplotlib.ticker.FuncFormatter(lambda v, _: f"{v:.1f}".replace(".", ","))

    # 1) tempo de parede por faixa de escopo
    faixas = [("focal\n(1 indicador)", lambda n: n == 1), ("intermediário\n(2 a 4)", lambda n: 2 <= n <= 4),
              ("amplo\n(todos os sistemas)", lambda n: n > 4)]
    fig, ax = plt.subplots(figsize=(6.3, 3.3))
    for i, (rot, cond) in enumerate(faixas):
        for j, (t, cor) in enumerate((("star", azul), ("hierarchical", laranja))):
            dados = [p[t]["wall_clock_ms"] / 1000 for c in Q if cond(n_ind[c]) for p in pares[c]]
            if not dados:
                continue
            bp = ax.boxplot([dados], positions=[i + (j - 0.5) * 0.34], widths=0.28, patch_artist=True,
                            showfliers=True, flierprops={"marker": "o", "markersize": 2.5, "markeredgecolor": cor, "alpha": 0.6},
                            medianprops={"color": tinta, "linewidth": 1.4}, whiskerprops={"color": cor},
                            capprops={"color": cor}, boxprops={"facecolor": cor, "edgecolor": cor, "alpha": 0.85})
            if t == "hierarchical":
                bp["boxes"][0].set_hatch("////")
                bp["boxes"][0].set_edgecolor("white")
    ax.set_xticks(range(len(faixas)))
    ax.set_xticklabels([f[0] for f in faixas])
    ax.set_ylabel("Tempo de parede (s)")
    ax.set_xlabel("Amplitude de escopo da pergunta")
    ax.grid(axis="x", visible=False)
    from matplotlib.patches import Patch
    ax.legend(handles=[Patch(facecolor=azul, label="Estrela"),
                       Patch(facecolor=laranja, hatch="////", edgecolor="white", label="Hierárquica")], loc="upper left")
    fig.savefig(os.path.join(pasta, "fig_tempo_por_escopo.png"))
    plt.close(fig)

    # 2) fidelidade: mediana por pergunta, estrela × hierárquica
    xs = [a["por_pergunta"][c]["ragas_faithfulness"]["star"] for c in Q]
    ys = [a["por_pergunta"][c]["ragas_faithfulness"]["hierarchical"] for c in Q]
    fig, ax = plt.subplots(figsize=(4.4, 4.2))
    lo, hi = 0.3, 1.0
    ax.fill_between([lo, hi], [lo - LIMIAR_INDIFERENCA, hi - LIMIAR_INDIFERENCA],
                    [lo + LIMIAR_INDIFERENCA, hi + LIMIAR_INDIFERENCA], color=grade, zorder=0, linewidth=0)
    ax.plot([lo, hi], [lo, hi], color=cinza, linewidth=0.8, linestyle="--", zorder=1)
    ax.scatter(xs, ys, s=22, color=tinta, alpha=0.75, linewidths=0.6, edgecolors="white", zorder=2)
    ax.set_xlim(lo, hi); ax.set_ylim(lo, hi); ax.set_aspect("equal")
    ax.xaxis.set_major_formatter(virgula); ax.yaxis.set_major_formatter(virgula)
    ax.set_xlabel("Fidelidade na estrela (mediana das rodadas)")
    ax.set_ylabel("Fidelidade na hierárquica (mediana das rodadas)")
    ax.text(0.33, 0.95, "hierárquica maior", color=cinza, fontsize=8, va="top")
    ax.text(0.97, 0.33, "estrela maior", color=cinza, fontsize=8, ha="right")
    fig.savefig(os.path.join(pasta, "fig_fidelidade_por_pergunta.png"))
    plt.close(fig)

    # 3) ruído: diferença absoluta de fidelidade entre topologias × entre repetições
    et = np.sort(a["ruido"]["ragas_faithfulness"]["_entre_topo"])
    er = np.sort(a["ruido"]["ragas_faithfulness"]["_entre_rep"])
    fig, ax = plt.subplots(figsize=(6.3, 3.3))
    if len(er):
        ax.step(er, np.arange(1, len(er) + 1) / len(er), where="post", color=cinza, linewidth=2, linestyle="--",
                label="Mesma topologia, rodadas diferentes")
    if len(et):
        ax.step(et, np.arange(1, len(et) + 1) / len(et), where="post", color=tinta, linewidth=2,
                label="Topologias diferentes, mesma rodada")
    ax.axvline(LIMIAR_INDIFERENCA, color=cinza, linewidth=0.8, linestyle=":")
    ax.text(LIMIAR_INDIFERENCA + 0.004, 0.04, "limiar de\nindiferença (0,05)", color=cinza, fontsize=8)
    ax.set_xlim(0, 0.5); ax.set_ylim(0, 1)
    ax.set_xlabel("Diferença absoluta de fidelidade")
    ax.set_ylabel("Proporção acumulada")
    ax.xaxis.set_major_formatter(virgula); ax.yaxis.set_major_formatter(virgula)
    ax.legend(loc="lower right")
    fig.savefig(os.path.join(pasta, "fig_ruido_da_fidelidade.png"))
    plt.close(fig)



def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--rodada", action="append", required=True,
                        help="jsonl de uma rodada (repita a opção; a ordem define o número da rodada)")
    parser.add_argument("--substituir-na-primeira", default=None,
                        help="jsonl cujos cenários substituem os da primeira rodada (reexecução após correção)")
    parser.add_argument("--out", default="analise_rodadas", help="pasta de saída")
    parser.add_argument("--sem-figuras", action="store_true")
    args = parser.parse_args()

    rodadas = [Rodada(i + 1, p) for i, p in enumerate(args.rodada)]
    if args.substituir_na_primeira:
        rodadas[0].substituir(args.substituir_na_primeira)
    a = analisar(rodadas)
    os.makedirs(args.out, exist_ok=True)
    if not args.sem_figuras:
        gravar_figuras(a, rodadas, args.out)
    limpo = _limpar(a)
    with open(os.path.join(args.out, "analise.json"), "w", encoding="utf-8") as f:
        json.dump(limpo, f, ensure_ascii=False, indent=2)
    gravar_xlsx(limpo, os.path.join(args.out, "analise.xlsx"))
    print(f"{a['rodadas']} rodada(s), {a['pareado']['perguntas_em_todas_as_rodadas']} perguntas pareadas, "
          f"{a['pareado']['analises']} análises -> {args.out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
