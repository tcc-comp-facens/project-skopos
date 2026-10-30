"""
Consolida as rodadas do experimento de revisão das entregas
(`batch_eval --experimento-revisao`) e produz as tabelas do eixo de
resiliência semântica.

Em cada análise do experimento, um agente de saúde entregou cerca de 20% dos
registros com valor errado (`scripts/falhas_semanticas.py`), e as duas
topologias revisaram as entregas pela mesma regra: o nó que recebe as
entregas dos agentes de dados as revisa numa chamada ao LLM e contesta os
valores suspeitos por reenvio (`agents/revisao_entregas.py`). O que muda entre
as topologias é onde a revisão fica: o orquestrador da estrela revisa tudo
numa chamada, e cada supervisor de dados da hierárquica revisa só o seu ramo.

Nada é recalculado a partir do sistema: o script só agrega as linhas que a
bateria gravou, já medidas contra o gabarito da corrupção.

Grandezas, por topologia e por amplitude de escopo da pergunta:

  sensibilidade      sinais em registro corrompido / registros corrompidos
  precisão           sinais em registro corrompido / sinais
  contaminados       registros que chegaram ao resultado ainda corrompidos
  perdidos           registros corrompidos descartados pela revisão
  achados limpos     análises cujos achados coincidem com os dos dados limpos
  custo              tokens, chamadas ao modelo, tempo de parede e de coordenação

Tratamento estatístico: teste exato de Fisher para a proporção de análises
com achados iguais aos dos dados limpos; Wilcoxon pareado (as duas topologias
de uma mesma análise receberam a mesma corrupção) para sensibilidade, tokens,
tempo de parede e custo de coordenação. Sem correção para comparações
múltiplas; todos os p são gravados.

Uso (na raiz do repositório; requer numpy, scipy e openpyxl):

    python backend/scripts/analise_revisao.py \
        --rodada revisao_rodada1.jsonl --rodada revisao_rodada2.jsonl \
        --out analise_revisao
"""

from __future__ import annotations

import argparse
import json
import os
import sys
from collections import defaultdict

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from scipy import stats  # noqa: E402

from scripts.analise_rodadas import NOME_TOPO, TOPOS, carregar, quartis, wilcoxon_pareado  # noqa: E402

FALHAS_SEMANTICAS = ("dados_corrompidos", "dados_corrompidos_persistente")
ESCOPOS = ("focal", "intermediario", "amplo")

# Grandezas comparadas por par (estrela, hierárquica) da mesma análise.
METRICAS_PAREADAS = [
    ("sensibilidade", "Sensibilidade da revisão"),
    ("total_tokens", "Tokens totais"),
    ("llm_call_count", "Chamadas ao modelo"),
    ("wall_clock_ms", "Tempo de parede (ms)"),
    ("overhead_ms", "Custo de coordenação (ms)"),
]


def _sensibilidade(row: dict) -> float | None:
    corrompidos = row.get("corrupcao_registros") or 0
    return (row.get("revisao_vp") or 0) / corrompidos if corrompidos else None


def coletar_pares(rodadas: list[list[dict]]) -> list[dict]:
    """Uma entrada por análise válida, com a linha de cada topologia.

    Válida: as duas topologias rodaram sem erro e o agente-alvo entregou dados
    corrompidos às duas (o alvo pode não ser acionado se a interpretação da
    pergunta mudar entre rodadas).
    """
    pares = []
    for numero, rows in enumerate(rodadas, start=1):
        por_cenario: dict[str, dict[str, dict]] = defaultdict(dict)
        for r in rows:
            if r.get("fault") in FALHAS_SEMANTICAS and r.get("topology") in TOPOS:
                por_cenario[r["scenario"]][r["topology"]] = r
        for cenario, linhas in por_cenario.items():
            if len(linhas) != 2 or any(l.get("run_error") for l in linhas.values()):
                continue
            if any(not (l.get("corrupcao_registros") or 0) for l in linhas.values()):
                continue
            pares.append({
                "rodada": numero, "cenario": cenario,
                "escopo": linhas["star"].get("escopo"), "alvo": linhas["star"].get("alvo_falha"),
                "fault": linhas["star"].get("fault"),
                **{t: dict(linhas[t], sensibilidade=_sensibilidade(linhas[t])) for t in TOPOS},
            })
    return pares


def _soma(pares: list[dict], topo: str, campo: str) -> int:
    return sum(p[topo].get(campo) or 0 for p in pares)


def resumir(pares: list[dict], topo: str) -> dict:
    corrompidos = _soma(pares, topo, "corrupcao_registros")
    vp = _soma(pares, topo, "revisao_vp")
    sinais = vp + _soma(pares, topo, "revisao_fp")
    limpos = [p[topo].get("achados_iguais_limpos") for p in pares if p[topo].get("achados_iguais_limpos") is not None]
    return {
        "analises": len(pares),
        "registros_corrompidos": corrompidos,
        "sinais": sinais,
        "sinais_em_corrompidos": vp,
        "sensibilidade": vp / corrompidos if corrompidos else None,
        "precisao": vp / sinais if sinais else None,
        "sensibilidade_por_analise": quartis([p[topo]["sensibilidade"] for p in pares]),
        "falsos_positivos_por_analise": quartis([p[topo].get("revisao_fp") or 0 for p in pares]),
        "corrigidos": _soma(pares, topo, "revisao_corrigidos"),
        "contaminados_finais": _soma(pares, topo, "contaminados_finais"),
        "analises_sem_contaminacao": sum(1 for p in pares if not p[topo].get("contaminados_finais")),
        "perdidos": _soma(pares, topo, "perdidos"),
        "achados_iguais_limpos": sum(limpos),
        "achados_avaliados": len(limpos),
        "revisao_falhou": _soma(pares, topo, "revisao_falhou"),
        "chamadas_de_revisao": quartis([p[topo].get("revisao_chamadas") or 0 for p in pares]),
        "total_tokens": quartis([p[topo].get("total_tokens") for p in pares]),
        "llm_call_count": quartis([p[topo].get("llm_call_count") for p in pares]),
        "wall_clock_ms": quartis([p[topo].get("wall_clock_ms") for p in pares]),
        "overhead_ms": quartis([p[topo].get("overhead_ms") for p in pares]),
    }


def comparar(pares: list[dict]) -> dict:
    """Testes entre as topologias sobre o mesmo conjunto de análises."""
    s = resumir(pares, "star")
    h = resumir(pares, "hierarchical")
    fisher = None
    if s["achados_avaliados"] and h["achados_avaliados"]:
        tabela = [
            [h["achados_iguais_limpos"], h["achados_avaliados"] - h["achados_iguais_limpos"]],
            [s["achados_iguais_limpos"], s["achados_avaliados"] - s["achados_iguais_limpos"]],
        ]
        fisher = float(stats.fisher_exact(tabela).pvalue)
    pareadas = {}
    for campo, rotulo in METRICAS_PAREADAS:
        difs = [p["hierarchical"].get(campo) - p["star"].get(campo) for p in pares
                if p["hierarchical"].get(campo) is not None and p["star"].get(campo) is not None]
        pareadas[campo] = {
            "rotulo": rotulo,
            "dif_mediana_hier_menos_estrela": quartis(difs).get("mediana"),
            "analises_hier_maior": sum(1 for d in difs if d > 0),
            "analises_hier_menor": sum(1 for d in difs if d < 0),
            "analises_iguais": sum(1 for d in difs if d == 0),
            "wilcoxon_p": wilcoxon_pareado(difs),
        }
    return {"fisher_achados_iguais_limpos_p": fisher, "pareadas": pareadas}


def analisar(rodadas: list[list[dict]], arquivos: list[str]) -> dict:
    pares = coletar_pares(rodadas)
    descartadas = sum(
        1 for rows in rodadas for r in rows
        if r.get("fault") in FALHAS_SEMANTICAS and r.get("topology") == "star"
    ) - len(pares)
    out: dict = {
        "rodadas": len(rodadas), "arquivos": arquivos,
        "falhas": sorted({p["fault"] for p in pares}),
        "analises_validas": len(pares), "analises_fora": descartadas,
        "geral": {t: resumir(pares, t) for t in TOPOS},
        "comparacao_geral": comparar(pares),
        "por_escopo": {},
        "por_analise": [
            {"rodada": p["rodada"], "cenario": p["cenario"], "escopo": p["escopo"], "alvo": p["alvo"],
             **{f"{campo}|{t}": p[t].get(campo) for t in TOPOS
                for campo in ("corrupcao_registros", "revisao_vp", "revisao_fp", "revisao_corrigidos",
                              "contaminados_finais", "perdidos", "achados_iguais_limpos",
                              "total_tokens", "wall_clock_ms")}}
            for p in pares
        ],
    }
    for escopo in ESCOPOS:
        sub = [p for p in pares if p["escopo"] == escopo]
        out["por_escopo"][escopo] = {
            **{t: resumir(sub, t) for t in TOPOS},
            "comparacao": comparar(sub) if sub else None,
        }
    return out


def gravar_xlsx(a: dict, path: str) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Resumo"
    ws.append(["Escopo", "Topologia", "Análises", "Registros corrompidos", "Sensibilidade", "Precisão",
               "Corrigidos", "Contaminados no resultado", "Perdidos", "Achados iguais aos limpos",
               "Revisão falhou", "Tokens (mediana)", "Tempo de parede ms (mediana)",
               "Custo de coordenação ms (mediana)", "Fisher p (achados limpos)"])
    blocos = [("todos", a["geral"], a["comparacao_geral"])] + [
        (e, {t: a["por_escopo"][e][t] for t in TOPOS}, a["por_escopo"][e]["comparacao"]) for e in ESCOPOS
    ]
    for escopo, resumo, comp in blocos:
        for t in TOPOS:
            r = resumo[t]
            ws.append([
                escopo, NOME_TOPO[t], r["analises"], r["registros_corrompidos"], r["sensibilidade"],
                r["precisao"], r["corrigidos"], r["contaminados_finais"], r["perdidos"],
                f'{r["achados_iguais_limpos"]} de {r["achados_avaliados"]}', r["revisao_falhou"],
                r["total_tokens"].get("mediana"), r["wall_clock_ms"].get("mediana"),
                r["overhead_ms"].get("mediana"), (comp or {}).get("fisher_achados_iguais_limpos_p"),
            ])
    ws = wb.create_sheet("Pareadas")
    ws.append(["Escopo", "Grandeza", "Diferença mediana (hier − estrela)", "Análises hier maior",
               "Análises hier menor", "Iguais", "Wilcoxon p"])
    for escopo, _, comp in blocos:
        for campo, d in ((comp or {}).get("pareadas") or {}).items():
            ws.append([escopo, d["rotulo"], d["dif_mediana_hier_menos_estrela"], d["analises_hier_maior"],
                       d["analises_hier_menor"], d["analises_iguais"], d["wilcoxon_p"]])
    ws = wb.create_sheet("PorAnalise")
    if a["por_analise"]:
        chaves = list(a["por_analise"][0])
        ws.append(chaves)
        for linha in a["por_analise"]:
            ws.append([linha.get(k) for k in chaves])
    wb.save(path)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--rodada", action="append", required=True,
                        help="jsonl de uma rodada do experimento (repita a opção)")
    parser.add_argument("--out", default="analise_revisao", help="pasta de saída")
    args = parser.parse_args()

    rodadas = [carregar(p) for p in args.rodada]
    a = analisar(rodadas, args.rodada)
    os.makedirs(args.out, exist_ok=True)
    with open(os.path.join(args.out, "analise_revisao.json"), "w", encoding="utf-8") as f:
        json.dump(a, f, ensure_ascii=False, indent=2)
    gravar_xlsx(a, os.path.join(args.out, "analise_revisao.xlsx"))
    print(f"{a['rodadas']} rodada(s), {a['analises_validas']} análises válidas "
          f"({a['analises_fora']} fora) -> {args.out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
